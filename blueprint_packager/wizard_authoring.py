"""Explicit, fail-closed Wizard bundle to Packager authoring conversion.

The extracted contracts are immutable inputs. This adapter does not reinterpret
the Wizard's review DOCX: it joins object identities in the structure and
activities contracts and writes the Packager's native authoring vocabulary.
Every conversion carries input hashes, a disposition ledger, and diagnostics.
"""
from __future__ import annotations

import hashlib
import json
import math
import os
import re
import tempfile
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, urlparse

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Pt, RGBColor

from . import vocabulary as V

SCHEMA = "blueprint_packager.wizard_authoring_conversion/1"
CONTRACTS = {
    "blueprint": ("__blueprint.json", {"coursecraft.blueprint/4", "coursecraft.blueprint/5"}),
    "activities": ("__course_activities.json", {"coursecraft.activities/1"}),
    "structure": ("__course_structure.json", {"coursecraft.structure/1"}),
    "rubrics": ("__rubrics.json", {"coursecraft.rubrics/1"}),
}


def _text(block: dict) -> str:
    return "".join(str(r.get("text", "")) for r in block.get("runs", []))


def _truth(value: Any) -> bool:
    return value is True or str(value).lower() in {"true", "1"}


def _number(value: Any) -> float | None:
    try:
        number = float(value)
        return number if math.isfinite(number) and number >= 0 else None
    except (TypeError, ValueError):
        return None


class _Conversion:
    def __init__(self, sources: dict, source_files: list[dict]) -> None:
        self.sources = sources
        self.report = {"schema": SCHEMA, "status": "blocked", "sources": source_files,
                       "diagnostics": [], "ledger": [], "output_path": None,
                       "extraction_diagnostics": {key: value.get("diagnostics", []) for key, value in sources.items()}}
        self.doc = Document()
        section = self.doc.sections[0]
        section.top_margin = section.bottom_margin = Inches(.65)
        section.left_margin = section.right_margin = Inches(.7)
        for name in ("Normal", "List Bullet", "List Number"):
            self.doc.styles[name].font.name = "Arial"
            self.doc.styles[name].font.size = Pt(11)
            self.doc.styles[name].paragraph_format.space_after = Pt(6)
        for name in ("Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
            self.doc.styles[name].font.name = "Arial"
            self.doc.styles[name].font.color.rgb = RGBColor.from_string("103945")
        for name in ("Title", "Subtitle"):
            for border in self.doc.styles[name].element.xpath("./w:pPr/w:pBdr"):
                border.getparent().remove(border)
        self.by_code: dict[str, tuple[str, dict]] = {}
        self.placements: dict[str, list[dict]] = {}
        self.page_targets: dict[str, str] = {}
        self.used_activities: set[str] = set()
        self.used_checklists: set[str] = set()
        self.used_pages: set[str] = set()
        self.rubric_users: dict[str, list[str]] = {}
        self.authored_text: set[str] = set()
        self.rubrics = {str(r.get("id")): r for r in sources["rubrics"].get("rubrics", [])}

    def note(self, code: str, message: str, where: str, severity: str = "error") -> None:
        self.report["diagnostics"].append(dict(severity=severity, code=code, message=message, where=where))

    def ledger(self, source: str, disposition: str, detail: str = "") -> None:
        self.report["ledger"].append(dict(source=source, disposition=disposition, detail=detail))

    def para(self, container, text: str = "", style: str | None = None):
        if hasattr(container, "_tc") and len(container.paragraphs) == 1 and not container.paragraphs[0].text:
            paragraph = container.paragraphs[0]
            paragraph.text = text
            if style:
                paragraph.style = style
        else:
            paragraph = container.add_paragraph(text, style=style)
        if text.startswith(("Page: ", "Page heading: ", "Assignment: ", "Discussion: ", "Checklist: ")):
            paragraph.paragraph_format.keep_with_next = True
        return paragraph

    def table(self, container, rows: int, cols: int):
        if hasattr(container, "_tc"):
            table = container.add_table(rows=rows, cols=cols, width=Inches(6))
        else:
            table = container.add_table(rows=rows, cols=cols)
        table.style = "Table Grid"
        for cell in table.rows[0].cells:
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "E8F1F4")
            cell._tc.get_or_add_tcPr().append(shading)
        header = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(header)
        return table

    def row_pair(self, table, label: str):
        header = table.add_row().cells[0]
        header.text = label
        header.paragraphs[0].runs[0].bold = True
        header.paragraphs[0].paragraph_format.keep_with_next = True
        shading = OxmlElement("w:shd"); shading.set(qn("w:fill"), "E8F1F4")
        header._tc.get_or_add_tcPr().append(shading)
        return table.add_row().cells[0]

    def link_run(self, paragraph, run: dict, where: str) -> None:
        text, href = str(run.get("text", "")), str(run.get("href") or "")
        if not text:
            return
        target = None
        if href and not urlparse(href).scheme:
            code = parse_qs(urlparse(href).query).get("rCode", [""])[0]
            if code in self.by_code:
                kind, obj = self.by_code[code]
                target = (kind, obj.get("name") or obj.get("title"))
            elif href in self.page_targets:
                target = ("page", self.page_targets[href])
            elif href.startswith("#"):
                self.note("LOCAL_ANCHOR_UNSUPPORTED", "A local anchor cannot be restored from extracted evidence", where)
            else:
                self.note("RELATIVE_LINK_UNRESOLVED", f"No exact extracted target for {href!r}", where)
            if target:
                if any(c in text for c in "|]") or any(c in target[1] for c in "|]"):
                    self.note("LINK_MARKER_UNREPRESENTABLE", "Link text or title contains reserved marker punctuation", where)
                text, href = f"[link to {target[0]}: {target[1]} | {text}]", ""
                self.ledger(where, "internal_link_to_authoring_marker", f"{target[0]}: {target[1]}")
        child = paragraph.add_run(text)
        child.bold = bool(run.get("bold"))
        child.italic = bool(run.get("italic"))
        child.underline = bool(run.get("underline"))
        if run.get("strike"):
            child.font.strike = True
        if run.get("sup") or run.get("sub"):
            self.note("RUN_FORMAT_UNSUPPORTED", "Superscript/subscript cannot be represented by the current authoring reader", where)
        if href:
            hyperlink = OxmlElement("w:hyperlink")
            rid = paragraph.part.relate_to(href, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
            hyperlink.set(qn("r:id"), rid)
            hyperlink.append(child._r)
            paragraph._p.append(hyperlink)

    def blocks(self, container, blocks: list, where: str) -> None:
        for index, block in enumerate(blocks):
            pointer = f"{where}/{index}"
            kind, text = block.get("kind"), _text(block)
            if kind == "divider" and not text.strip():
                self.ledger(pointer, "layout_only", "Empty horizontal divider")
                continue
            if text:
                self.authored_text.add(text)
            if (kind == "p" and text == "↑ Back to top" and len(block.get("runs", [])) == 1
                    and block["runs"][0].get("href") == "#top"):
                self.ledger(pointer, "generated_navigation", "The page template regenerates Back to top")
                continue
            if kind in {"dropdown", "visual"}:
                key = V.D_DROPDOWN if kind == "dropdown" else V.D_VISUAL
                end = V.D_DROPDOWN_END if kind == "dropdown" else V.D_VISUAL_END
                meta = block.get("meta", {})
                attrs = [(key, meta[key]) for key in ("tone", "color") if meta.get(key)]
                if kind == "visual" and attrs:
                    # These are explicit extracted semantic hints, not a CSS
                    # inference. The authoring reader carries bracket attributes.
                    tag = {"Callout": "callout", "Styled section": "box", "Timeline": "timeline"}.get(text)
                    if tag is None or any(re.search(r"\s|[\[\]]", str(value)) for _, value in attrs):
                        self.note("VISUAL_ATTRIBUTES_UNSUPPORTED", "Visual attributes cannot be represented by the authoring vocabulary", pointer)
                        continue
                    self.para(container, "[" + tag + " " + " ".join(f"{k}={v}" for k, v in attrs) + "]")
                else:
                    self.para(container, f"{key}: {text}")
                self.blocks(container, block.get("blocks", []), pointer + "/blocks")
                self.para(container, end)
            elif kind == "table":
                rows = block.get("rows", [])
                if not rows or not all(len(row) == len(rows[0]) for row in rows):
                    self.note("TABLE_SHAPE", "Empty or irregular extracted table", pointer)
                    continue
                table = self.table(container, len(rows), len(rows[0]))
                for ri, row in enumerate(rows):
                    for ci, runs in enumerate(row):
                        para = table.cell(ri, ci).paragraphs[0]
                        for run in runs:
                            self.link_run(para, run, f"{pointer}/rows/{ri}/{ci}")
            elif kind == "embed":
                runs = block.get("runs", [])
                if len(runs) != 1 or not runs[0].get("href") or block.get("meta", {}).get("embed_type") != "Video embed":
                    self.note("EMBED_UNSUPPORTED", "Only a single extracted video embed is supported", pointer)
                    continue
                para = self.para(container)
                self.link_run(para, {"text": f"{V.D_VIDEO}: {text}", "href": runs[0]["href"]}, pointer)
            elif kind in {"p", "li", "label"}:
                level = int(block.get("level") or 0)
                style = f"Heading {max(2, min(6, level))}" if kind == "label" and level else None
                if kind == "li":
                    style = "List Number" if block.get("meta", {}).get("list") == "ordered" else "List Bullet"
                para = self.para(container, style=style)
                if kind == "li":
                    # Explicit numPr survives Word and Google conversion; the reader
                    # also understands style-linked definitions.
                    style_pr = self.doc.styles[style].element.find(qn("w:pPr"))
                    numpr = style_pr.find(qn("w:numPr")) if style_pr is not None else None
                    numid = numpr.find(qn("w:numId")) if numpr is not None else None
                    if numid is not None:
                        pr = para._p.get_or_add_pPr().get_or_add_numPr()
                        pr.get_or_add_numId().val = int(numid.get(qn("w:val")))
                        pr.get_or_add_ilvl().val = max(0, level - 1)
                for run in block.get("runs", []):
                    self.link_run(para, run, pointer)
            else:
                self.note("BLOCK_UNSUPPORTED", f"Extracted block kind {kind!r} needs an explicit conversion rule", pointer)
                continue
            self.ledger(pointer, "authored_content", kind)

    def prepare(self) -> None:
        activities = self.sources["activities"]
        if len(self.rubrics) != len(self.sources["rubrics"].get("rubrics", [])):
            self.note("RUBRIC_IDENTITY", "Rubric ids must be unique", "rubrics")
        for group, kind in (("dropbox_folders", "assignment"), ("discussions", "discussion"), ("checklists", "checklist")):
            for index, obj in enumerate(activities.get(group, [])):
                if kind == "discussion" and obj.get("kind") == "forum":
                    if obj.get("description_text") or obj.get("description_html"):
                        self.note("FORUM_CONTENT_UNSUPPORTED", "Forum descriptions have no authoring field", f"activities/{group}/{index}")
                    if _truth(obj.get("is_hidden")):
                        self.note("FORUM_VISIBILITY_UNSUPPORTED", "Forum visibility has no authoring field", f"activities/{group}/{index}")
                    continue
                code = obj.get("resource_code", "")
                if not code or code in self.by_code:
                    self.note("ACTIVITY_IDENTITY", "Missing or duplicate activity resource code", f"activities/{group}/{index}")
                self.by_code[code] = kind, obj
                for field in ("date_due", "date_start", "date_end", "condition_set"):
                    if obj.get(field):
                        self.note("ACTIVITY_SETTINGS_UNSUPPORTED", f"{field} requires separate explicit build settings; conversion would lose it", f"activities/{group}/{index}")
                if kind == "assignment" and (obj.get("submission_type", "0") != "0" or obj.get("folder_type", "2") != "2"):
                    self.note("SUBMISSION_TYPE_UNSUPPORTED", "Only the Packager's individual file submission shape is supported", f"activities/{group}/{index}")
                if kind == "discussion" and any(_truth(obj.get(k)) for k in ("requires_approval", "must_post_to_participate")):
                    self.note("DISCUSSION_SETTINGS_UNSUPPORTED", "Moderation/post-first settings need an explicit authoring representation", f"activities/{group}/{index}")
        if activities.get("quizzes"):
            self.note("QUIZ_SOURCE_REQUIRED", "Quiz carry-forward needs the original package and an explicit Source directive; conversion cannot invent it", "activities/quizzes")
        if activities.get("condition_sets"):
            self.note("CONDITIONS_UNSUPPORTED", "Release conditions cannot be represented by this converter", "activities/condition_sets")
        for root in self.sources["structure"].get("tree", []):
            if root.get("kind") != "module":
                self.note("ROOT_KIND_UNSUPPORTED", "The authoring model requires top-level modules", "structure/tree")
                continue
            if root.get("description"):
                self.note("MODULE_DESCRIPTION_UNSUPPORTED", "Module description has no authoring field", root["title"])
            if V.WEEK_HEADING_RE.match(root.get("title", "")) and _truth(root.get("is_hidden")):
                self.note("WEEK_VISIBILITY_UNSUPPORTED", "Hidden weekly modules have no authoring visibility field", root["title"])
            for node in root.get("children", []):
                if node.get("children"):
                    self.note("NESTED_MODULE_UNSUPPORTED", "Nested modules need an explicit structure conversion decision", node.get("title", ""))
                code = node.get("rcode")
                if code:
                    self.placements.setdefault(code, []).append(root)
                    if code in self.by_code:
                        _kind, target = self.by_code[code]
                        if node.get("title") != (target.get("name") or target.get("title")):
                            self.note("MANIFEST_TITLE_DIFFERS", "A manifest alias differs from the activity title and needs an explicit representation", node.get("title", ""))
                        if _truth(node.get("is_hidden")) != _truth(target.get("is_hidden")):
                            self.note("PLACEMENT_VISIBILITY_DIFFERS", "Manifest and activity visibility differ", node.get("title", ""))
                if node.get("kind") == "html_topic":
                    href = node.get("href", "")
                    if href in self.page_targets and self.page_targets[href] != node["title"]:
                        self.note("PAGE_TITLE_ALIAS_UNSUPPORTED", "One page address has multiple manifest titles", href)
                    self.page_targets[href] = node["title"]
                    base = Path(href).name
                    if base in self.page_targets and self.page_targets[base] != node["title"]:
                        self.note("PAGE_BASENAME_COLLISION", f"Ambiguous relative page address {base}", node["title"])
                    self.page_targets[base] = node["title"]
        for code, placements in self.placements.items():
            if len(placements) > 1:
                self.note("ACTIVITY_MULTIPLE_PLACEMENTS", "Multiple activity/checklist placements need an explicit conversion representation", code)
        # Contract identity prevents accidental mixing of separate extraction runs.
        values = {self.sources[k].get("run_id") for k in ("activities", "structure") if self.sources[k].get("run_id")}
        if len(values) > 1:
            self.note("BUNDLE_RUN_MISMATCH", "Activities and structure come from different extraction runs", "bundle")

    def activity(self, container, code: str, *, weekly: bool = False) -> None:
        if code in self.used_activities:
            self.note("ACTIVITY_MULTIPLE_PLACEMENTS", "Repeated activity placements are not representable in one authoring block", code)
            return
        self.used_activities.add(code)
        kind, obj = self.by_code[code]
        title = obj.get("name") if kind == "assignment" else obj.get("title")
        where = f"activities/{kind}/{code}"
        self.para(container, f"{kind.title()}: {title}")
        points = obj.get("out_of") if kind == "assignment" else obj.get("score_out_of")
        if points not in (None, ""):
            number = _number(points)
            if number is None:
                self.note("ACTIVITY_POINTS", f"Unreadable points {points!r}", where)
            else:
                self.para(container, f"Points: {number:g}")
        if obj.get("grade_item_code"):
            grades = [g for g in self.sources["activities"].get("grade_items", []) if g.get("resource_code") == obj["grade_item_code"]]
            if len(grades) != 1:
                self.note("GRADE_JOIN_UNRESOLVED", "Grade item code must resolve exactly once", where)
            else:
                self.para(container, f"Gradebook item: {grades[0]['name']}")
        rids = obj.get("rubric_ids", [])
        if len(rids) > 1:
            self.note("MULTIPLE_RUBRICS_UNSUPPORTED", "The authoring model supports one rubric per activity", where)
        for rid in rids:
            rubric = self.rubrics.get(str(rid))
            if rubric is None:
                self.note("RUBRIC_JOIN_UNRESOLVED", f"Unknown rubric id {rid}", where)
            else:
                self.para(container, f"Rubric: {rubric['name']}")
                self.rubric_users.setdefault(str(rid), []).append(f"{kind.title()}: {title}")
        if kind == "discussion":
            forums = [d for d in self.sources["activities"].get("discussions", []) if d.get("kind") == "forum" and d.get("forum_id") == obj.get("forum_id")]
            if len(forums) != 1:
                self.note("FORUM_JOIN_UNRESOLVED", "Forum id must resolve exactly once", where)
            else:
                self.para(container, f"Forum: {forums[0]['title']}")
        if _truth(obj.get("is_hidden")):
            self.para(container, f"Visibility: {V.VISIBILITY_HIDDEN}")
        if not weekly:
            for parent in self.placements.get(code, []):
                self.para(container, f"Placement: {parent['title']}")
        key = "instructions_blocks" if kind == "assignment" else "description_blocks"
        if key not in obj and (obj.get("instructions_text") or obj.get("description_text")):
            self.note("ACTIVITY_BLOCKS_MISSING", "Only flattened text is available; structured authored content is required", where)
        self.blocks(container, obj.get(key, []), where + "/" + key)
        self.ledger(where, "activity", title)

    def page(self, container, node: dict, parent: dict) -> None:
        href = node.get("href", "")
        if href in self.used_pages:
            return
        self.used_pages.add(href)
        pages = [p for p in self.sources["structure"].get("html_topics", []) if p.get("href") == href and not p.get("hidden_manifest_item")]
        if len(pages) != 1:
            self.note("PAGE_EVIDENCE_MISSING", "HTML topic must resolve to one extracted page body", node["title"])
            return
        page = pages[0]
        self.para(container, f"Page: {node['title']}")
        if page.get("html_title"):
            self.para(container, f"Page heading: {page['html_title']}")
            self.authored_text.add(page["html_title"])
        for root in self.sources["structure"].get("tree", []):
            if any(c.get("href") == href and c.get("kind") == "html_topic" for c in root.get("children", [])):
                self.para(container, f"Placement: {root['title']}")
        if _truth(node.get("is_hidden")):
            self.para(container, f"Visibility: {V.VISIBILITY_HIDDEN}")
        for i, segment in enumerate(page.get("body_segments", [])):
            if segment.get("heading") and not (i == 0 and segment.get("level") == 1):
                self.para(container, segment["heading"], f"Heading {max(2, min(6, segment.get('level') or 2))}")
                self.authored_text.add(segment["heading"])
            self.blocks(container, segment.get("blocks", []), f"structure/html_topics/{href}/body_segments/{i}/blocks")
        self.ledger(f"structure/html_topics/{href}", "page", node["title"])

    def grade_plan(self) -> None:
        rows = self.sources["activities"].get("grade_items", [])
        table = self.table(self.doc, len(rows) + 2, 2)
        table.autofit = False
        for row in table.rows:
            row.cells[0].width, row.cells[1].width = Inches(5.4), Inches(1.7)
        table.columns[0].width, table.columns[1].width = Inches(5.4), Inches(1.7)
        table.cell(0, 0).text, table.cell(0, 1).text = V.GRADE_PLAN_HEADER
        total = 0.
        for index, row in enumerate(rows, 1):
            where = f"activities/grade_items/{index - 1}"
            points = _number(row.get("out_of"))
            if points is None:
                self.note("GRADE_POINTS_MISSING", "Numeric grade maximum is required", where)
                points = 0.
            if str(row.get("type_id", "1")) != "1" or row.get("category_id") or _truth(row.get("is_bonus")) or _truth(row.get("exclude_from_final_grade_calc")):
                self.note("GRADE_SETTINGS_UNSUPPORTED", "Only uncategorized ordinary numeric grades are supported", where)
            if row.get("is_active") not in (None, "") and not _truth(row["is_active"]):
                self.note("GRADE_SETTINGS_UNSUPPORTED", "Inactive grade items have no authoring representation", where)
            total += points
            table.cell(index, 0).text = row.get("name", "")
            table.cell(index, 1).text = f"{points:g}"
            self.ledger(where, "grade_plan_row", row.get("name", ""))
        table.cell(len(rows) + 1, 0).text = V.GRADE_PLAN_TOTAL
        table.cell(len(rows) + 1, 1).text = f"{total:g}"

    def week(self, week: dict, module: dict) -> None:
        title = week["title"]
        self.para(self.doc, title, "Heading 1")
        table = self.doc.add_table(rows=0, cols=1)
        table.style = "Table Grid"
        children = module.get("children", [])
        pages = [n for n in children if n.get("kind") == "html_topic"]
        heading = _text(week.get("overview", [{}])[0]).removesuffix(":") if week.get("overview") else ""
        candidates = [p for p in pages if any(t.get("href") == p.get("href") and t.get("html_title") == heading for t in self.sources["structure"].get("html_topics", []))]
        if len(candidates) != 1:
            self.note("OVERVIEW_JOIN_UNRESOLVED", "Overview heading must match exactly one extracted HTML page title", title)
            return
        overview = candidates[0]
        if _truth(overview.get("is_hidden")):
            self.note("OVERVIEW_VISIBILITY_UNSUPPORTED", "Hidden overview pages have no authoring visibility field", title)
        self.used_pages.add(overview["href"])
        for key, _prefix, label in V.SECTION_LABELS:
            cell = self.row_pair(table, label)
            if key == "overview":
                self.para(cell, f"Page: {overview['title']}")
                self.para(cell, f"Page heading: {heading}")
                self.blocks(cell, week.get("overview", [])[1:], f"blueprint/weeks/{title}/overview/1:")
                self.ledger(f"blueprint/weeks/{title}/overview/0", "page_heading", heading)
            elif key == "objectives":
                self.blocks(cell, week.get("learning_objectives", []), f"blueprint/weeks/{title}/learning_objectives")
            elif key == "resources":
                for index, section in enumerate(week.get("resources", [])):
                    if section.get("label"):
                        self.para(cell, section["label"], f"Heading {max(2, min(6, section.get('level') or 2))}")
                    self.blocks(cell, section.get("blocks", []), f"blueprint/weeks/{title}/resources/{index}/blocks")
            elif key in {"assignments", "discussions"}:
                kind = "assignment" if key == "assignments" else "discussion"
                for node in children:
                    code = node.get("rcode")
                    if code in self.by_code and self.by_code[code][0] == kind:
                        self.activity(cell, code, weekly=True)
            elif key == "checklist":
                checks = [n for n in children if n.get("kind") == "checklist_link"]
                if len(checks) > 1:
                    self.note("MULTIPLE_CHECKLISTS_UNSUPPORTED", "Only one checklist per week is supported", title)
                for node in checks:
                    code = node.get("rcode")
                    if code not in self.by_code or self.by_code[code][0] != "checklist":
                        self.note("CHECKLIST_JOIN_UNRESOLVED", "Checklist resource code is missing from activities", title)
                        continue
                    self.used_checklists.add(code)
                    checklist = self.by_code[code][1]
                    self.para(cell, f"Checklist: {checklist['name']}")
                    blocks = checklist.get("blocks", [])
                    if blocks and blocks[0].get("kind") == "label" and _text(blocks[0]) == checklist["name"] + ":":
                        self.ledger(f"activities/checklists/{code}/blocks/0", "checklist_title", checklist["name"])
                        blocks = blocks[1:]
                    if any(b.get("kind") != "li" for b in blocks):
                        self.note("CHECKLIST_CATEGORIES_UNSUPPORTED", "Checklist category structure needs an explicit conversion rule", title)
                    self.blocks(cell, blocks, f"activities/checklists/{code}/blocks")
        # Extra pages remain explicit course-level Page objects with a Placement
        # directive; they are written in the introduction section by render().

    def rubric(self, rid: str, rubric: dict) -> None:
        where = f"rubrics/{rid}"
        self.para(self.doc, rubric["name"], "Heading 2")
        for user in self.rubric_users.get(rid, []):
            self.para(self.doc, f"Used by: {user}")
        if rubric.get("description"):
            self.para(self.doc, rubric["description"])
        levels, criteria = rubric.get("levels", []), rubric.get("criteria", [])
        expected_attributes = {"type": "1", "scoring_method": "3", "display_levels_in_des_order": "True",
                               "state": "0", "visibility": "0", "uses_overall_score": "True",
                               "has_manual_alignment": "False", "score_visible_to_assessed_users": "True",
                               "enabled_feedback_copy": "False", "usage_restrictions": "Competency,ePortfolio"}
        for key, value in rubric.get("attributes", {}).items():
            if key in expected_attributes and str(value).lower() != expected_attributes[key].lower():
                self.note("RUBRIC_SETTINGS_UNSUPPORTED", f"Rubric {key}={value!r} differs from the supported Packager shape", where)
        if rubric.get("scoring_method") != "3" or not levels or not criteria:
            self.note("RUBRIC_SHAPE_UNSUPPORTED", "Only complete analytic custom-points rubrics are supported", where)
            return
        table = self.table(self.doc, len(criteria) + 1, len(levels) + 1)
        table.cell(0, 0).text = "Criteria"
        percentages = []
        for i, level in enumerate(levels, 1):
            band = str(level.get("score_band", ""))
            match = re.fullmatch(r"(\d+(?:\.\d+)?)%\+?", band)
            pct = float(match[1]) if match else None
            if pct is None:
                self.note("RUBRIC_LEVEL_UNSUPPORTED", "A percentage score band is required", where)
                pct = 0.
            percentages.append(pct)
            table.cell(0, i).text = f"{level['name']} — {pct:g}%"
        for row, criterion in enumerate(criteria, 1):
            name = criterion.get("name", "")
            weight_match = re.search(r"(\d+(?:\.\d+)?)%\)?\s*$", name)
            if not weight_match:
                self.note("RUBRIC_WEIGHT_MISSING", "Criterion name does not state its percentage weight", where)
            weight = float(weight_match[1]) if weight_match else 0.
            table.cell(row, 0).text = name
            cells = {str(c.get("level_id")): c for c in criterion.get("cells", [])}
            for col, (level, pct) in enumerate(zip(levels, percentages), 1):
                cell = cells.get(str(level["level_id"]))
                if cell is None or _number(cell.get("points")) is None or abs(float(cell["points"]) - weight * pct / 100) > 1e-6:
                    self.note("RUBRIC_SCORE_UNREPRESENTABLE", "Cell points disagree with criterion weight and level percentage", where)
                table.cell(row, col).text = (cell or {}).get("description", "")
        if any(level.get("description") or level.get("feedback") for level in rubric.get("overall_levels", [])):
            self.note("RUBRIC_OVERALL_TEXT_UNSUPPORTED", "Overall-level descriptions/feedback have no authoring field", where)
        self.ledger(where, "rubric_grid", rubric["name"])

    def render(self) -> None:
        self.prepare()
        blueprint, structure = self.sources["blueprint"], self.sources["structure"]
        self.para(self.doc, "Blueprint", "Title")
        self.para(self.doc, " ".join(str(blueprint.get(k) or "") for k in ("course_number", "course_title", "term")).strip())
        for key, label in V.FRONT_LABELS.items():
            blocks = blueprint.get("front_matter", {}).get(key, [])
            if blocks:
                table = self.table(self.doc, 2, 1)
                table.cell(0, 0).text = label
                self.blocks(table.cell(1, 0), blocks, f"blueprint/front_matter/{key}")
            else:
                self.note("FRONT_MATTER_NOT_EXTRACTED", f"{label} is absent from extracted evidence; no content was invented", f"blueprint/front_matter/{key}", "warning")
        self.grade_plan()
        weeks = blueprint.get("weeks", [])
        modules = structure.get("tree", [])
        week_modules = {}
        for week in weeks:
            matches = [m for m in modules if m.get("title") == week.get("title") and m.get("kind") == "module"]
            if len(matches) != 1 or not V.WEEK_HEADING_RE.match(week.get("title", "")):
                self.note("WEEK_JOIN_UNRESOLVED", "Week title must match exactly one top-level module and use the Week N - Topic frame", week.get("title", ""))
            else:
                week_modules[week["title"]] = matches[0]
        for module in modules:
            if V.WEEK_HEADING_RE.match(module.get("title", "")) and module["title"] not in week_modules:
                self.note("WEEK_EVIDENCE_MISSING", "A weekly manifest module is absent from the extracted Blueprint weeks", module["title"])
        # Generate weeks in a detached document to discover overview ownership
        # before writing course pages, without guessing by title substring.
        front_doc = self.doc
        self.doc = Document()
        # The temporary week document supplies relationships while assembling
        # sections. Match the final page geometry so its table grid widths do
        # not retain Word's narrower default margin layout when moved.
        for attr in ("page_width", "page_height", "left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(self.doc.sections[0], attr, getattr(front_doc.sections[0], attr))
        for week in weeks:
            if week["title"] in week_modules:
                self.week(week, week_modules[week["title"]])
        week_doc = self.doc
        self.doc = front_doc
        self.para(self.doc, V.H_COURSE_INTRO, "Heading 1")
        for module in modules:
            for node in module.get("children", []):
                if node.get("kind") == "html_topic" and node.get("href") not in self.used_pages:
                    self.page(self.doc, node, module)
        for page in structure.get("html_topics", []):
            if not page.get("hidden_manifest_item") and page.get("href") not in self.used_pages:
                self.note("PAGE_UNPLACED", "An extracted HTML topic has no represented manifest placement", page.get("href", ""))
        self.para(self.doc, V.H_BEFORE_WEEK_1, "Heading 1")
        table = self.doc.add_table(rows=0, cols=1); table.style = "Table Grid"
        for module in modules:
            if module.get("title") not in week_modules:
                cell = self.row_pair(table, module.get("title", ""))
                self.para(cell, f"Module: {module.get('title', '')}")
                if _truth(module.get("is_hidden")):
                    self.para(cell, f"Visibility: {V.VISIBILITY_HIDDEN}")
            for node in module.get("children", []):
                kind = node.get("kind")
                if kind == "quicklink":
                    cell = self.row_pair(table, node.get("title", ""))
                    self.para(cell, f"Link: {node['title']}")
                    self.para(cell, f"URL: {node.get('href', '')}")
                    self.para(cell, f"Placement: {module['title']}")
                    if _truth(node.get("is_hidden")):
                        self.para(cell, f"Visibility: {V.VISIBILITY_HIDDEN}")
                    if node.get("description"):
                        self.note("LINK_DESCRIPTION_UNSUPPORTED", "Manifest link description has no authoring field", node["title"])
                    self.ledger(f"structure/tree/{node.get('identifier')}", "external_link", node["title"])
                elif kind not in {"html_topic", "dropbox_link", "discussion_link", "checklist_link"}:
                    self.note("MANIFEST_KIND_UNSUPPORTED", f"Unsupported manifest object kind {kind!r}", node.get("title", ""))
                elif kind.endswith("_link") and node.get("rcode") not in self.by_code:
                    self.note("MANIFEST_JOIN_UNRESOLVED", "Manifest resource code has no extracted activity", node.get("title", ""))
        # Move week XML and its hyperlinks into the final document. Relationship
        # identifiers belong to a document part and must be copied explicitly.
        for element in list(week_doc.element.body):
            if element.tag == qn("w:sectPr"):
                continue
            for hyperlink in element.iter(qn("w:hyperlink")):
                old = hyperlink.get(qn("r:id"))
                if old:
                    rel = week_doc.part.rels[old]
                    hyperlink.set(qn("r:id"), self.doc.part.relate_to(rel.target_ref, rel.reltype, is_external=rel.is_external))
            self.doc.element.body.insert(len(self.doc.element.body) - 1, element)
        self.para(self.doc, V.H_COURSE_ACTIVITIES, "Heading 1")
        for code, (kind, obj) in self.by_code.items():
            if kind in {"assignment", "discussion"} and code not in self.used_activities:
                self.activity(self.doc, code)
            elif kind == "checklist" and code not in self.used_checklists:
                self.note("UNPLACED_CHECKLIST_UNSUPPORTED", "A checklist must belong to exactly one extracted week", obj.get("name", ""))
        self.para(self.doc, V.H_RUBRICS, "Heading 1")
        for rid, rubric in self.rubrics.items():
            self.rubric(rid, rubric)
        # The Blueprint contract repeats content found in the structure and
        # activities contracts. Verify those alternate views instead of silently
        # ignoring them; an unmatched authored passage means a mixed or lossy
        # bundle and must stop conversion.
        def audit_aliases(blocks: list, where: str) -> None:
            for i, block in enumerate(blocks):
                pointer = f"{where}/{i}"
                kind, text = block.get("kind"), _text(block)
                if kind == "hidden":
                    self.ledger(pointer, "manifest_evidence", "Visibility and identity are read from structure/tree")
                elif kind == "divider" and not text:
                    self.ledger(pointer, "layout_only", "Empty horizontal divider")
                elif kind == "label" and text.startswith(("Points: ", "Rubric: ", "Gradebook item: ")):
                    self.ledger(pointer, "activity_metadata", "Resolved through activities object identities")
                elif kind == "label" and not block.get("level"):
                    # Checklist's tool label is generated, but any category
                    # label is rejected by week() above.
                    self.ledger(pointer, "checklist_label", text)
                elif kind == "label" and text.removesuffix(":") in self.authored_text:
                    self.ledger(pointer, "heading_alias", "Heading carried from structure evidence; Wizard label punctuation is structural")
                elif text and text not in self.authored_text:
                    self.note("BLUEPRINT_CONTENT_UNACCOUNTED", f"Authored text in the Blueprint view is absent from converted object evidence: {text[:100]!r}", pointer)
                else:
                    self.ledger(pointer, "duplicate_extracted_view", "Content represented by object evidence")
                audit_aliases(block.get("blocks", []), pointer + "/blocks")
        audit_aliases(blueprint.get("front_matter", {}).get("course_introduction", []), "blueprint/front_matter/course_introduction")
        for i, section in enumerate(blueprint.get("before_week_1", [])):
            audit_aliases(section.get("blocks", []), f"blueprint/before_week_1/{i}/blocks")
        for i, week in enumerate(weeks):
            for key in ("assignments", "discussions", "checklist", "other_sections"):
                for j, section in enumerate(week.get(key, [])):
                    audit_aliases(section.get("blocks", []), f"blueprint/weeks/{i}/{key}/{j}/blocks")
        for key, sections in blueprint.get("unplaced_activities", {}).items():
            for i, section in enumerate(sections):
                audit_aliases(section.get("blocks", []), f"blueprint/unplaced_activities/{key}/{i}/blocks")
        self.report["limitations"] = [
            "Local authoring conversion only; no tenant import or playback is implied.",
            "Dates, release conditions, quizzes, nested modules, and unsupported authored blocks stop conversion.",
            "Decorative dividers and exact generated Back to top links are recorded as layout/navigation.",
            "Extraction may not recover unpublished front matter, source formatting, or settings absent from its contracts.",
        ]


def convert_bundle(bundle_dir: Path, output_path: Path, *, force: bool = False) -> dict:
    """Convert one Wizard bundle, writing ``<stem>.conversion.json`` alongside it.

    Returns a reviewable result with ``status`` equal to ``converted`` or
    ``blocked``. Errors produce a ledger but no DOCX. Existing output/ledger
    paths raise FileExistsError unless ``force`` is explicit. Input contracts
    are never modified. Bundle discovery requires one file per applicable role;
    absent rubric evidence needs an inventory proving a rubric-free package.
    """
    bundle_dir, output_path = Path(bundle_dir), Path(output_path)
    ledger_path = output_path.with_suffix(".conversion.json")
    if output_path.suffix.lower() != ".docx":
        raise ValueError("Authoring output must have a .docx extension")
    for path in (output_path, ledger_path):
        if path.exists() and not force:
            raise FileExistsError(f"Refusing to replace existing file: {path}")
        if path.resolve().is_relative_to(bundle_dir.resolve()):
            raise ValueError("Write authoring output outside the immutable extraction bundle")
    sources, files, errors = {}, [], []
    rubric_contract_absent = False
    for role, (suffix, schemas) in CONTRACTS.items():
        candidates = list(bundle_dir.glob("*" + suffix))
        if role == "rubrics" and not candidates:
            rubric_contract_absent = True
            continue
        if len(candidates) != 1:
            errors.append(dict(severity="error", code="BUNDLE_FILE_COUNT", message=f"Expected one {suffix}; found {len(candidates)}", where=role))
            continue
        path = candidates[0]
        data = path.read_bytes()
        files.append(dict(role=role, path=str(path.resolve()), sha256=hashlib.sha256(data).hexdigest(), bytes=len(data)))
        try:
            value = json.loads(data)
        except (ValueError, UnicodeDecodeError):
            errors.append(dict(severity="error", code="BUNDLE_JSON", message="Invalid JSON", where=role))
            continue
        required = {
            "blueprint": ("course_title", "front_matter", "before_week_1", "weeks"),
            "activities": ("dropbox_folders", "discussions", "checklists", "grade_items"),
            "structure": ("tree", "html_topics"),
            "rubrics": ("rubrics",),
        }
        if not isinstance(value, dict) or value.get("schema") not in schemas:
            errors.append(dict(severity="error", code="BUNDLE_SCHEMA", message=f"Expected one of {sorted(schemas)}", where=role))
        elif any(key not in value for key in required[role]):
            errors.append(dict(severity="error", code="BUNDLE_SHAPE_INVALID", message=f"Required fields: {', '.join(required[role])}", where=role))
        else:
            sources[role] = value
    if rubric_contract_absent:
        activity_data = sources.get("activities", {})
        rows = activity_data.get("dropbox_folders", []) + activity_data.get("discussions", []) + activity_data.get("quizzes", [])
        has_refs = (any(row.get("rubric_ids") or row.get("rubrics_resolved") for row in rows)
                    or any(j.get("join_type") == "rubric" for j in activity_data.get("joins", [])))
        inventories = list(bundle_dir.glob("*__inventory.json"))
        inventory = None
        if len(inventories) == 1:
            data = inventories[0].read_bytes()
            try:
                inventory = json.loads(data)
            except (ValueError, UnicodeDecodeError):
                pass
            files.append(dict(role="inventory", path=str(inventories[0].resolve()),
                              sha256=hashlib.sha256(data).hexdigest(), bytes=len(data)))
        names = inventory.get("all_files") if isinstance(inventory, dict) else None
        proven_absent = (isinstance(names, list) and all(isinstance(name, str) for name in names)
                         and not any("rubric" in Path(name).name.lower() for name in names))
        if has_refs or not proven_absent or "activities" not in sources:
            errors.append(dict(severity="error", code="RUBRIC_CONTRACT_MISSING",
                               message="Absent rubrics JSON requires zero rubric associations and an inventory proving no rubric payload exists", where="rubrics"))
        else:
            sources["rubrics"] = {"rubrics": [], "diagnostics": []}
            files.append(dict(role="rubrics", path=None, status="not_applicable",
                              reason="No extracted rubric associations; inventory lists no rubric payload"))
    candidate = None
    if errors:
        report = dict(schema=SCHEMA, status="blocked", sources=files, diagnostics=errors, ledger=[], output_path=None)
    else:
        conversion = _Conversion(sources, files)
        try:
            conversion.render()
        except (KeyError, TypeError, ValueError, IndexError) as exc:
            conversion.note("BUNDLE_SHAPE_INVALID", f"Contract data is incomplete or malformed: {exc}", "bundle")
        report = conversion.report
        if not any(d["severity"] == "error" for d in report["diagnostics"]):
            output_path.parent.mkdir(parents=True, exist_ok=True)
            with tempfile.NamedTemporaryFile(prefix=".wizard-authoring-", suffix=".docx", dir=output_path.parent, delete=False) as temporary:
                candidate = Path(temporary.name)
            try:
                conversion.doc.save(candidate)
                # Read back before publishing either output.
                from .reader import read_model
                from .joins import resolve
                model = read_model(candidate)
                diagnostics = model.get("diagnostics", []) + resolve(model)["diagnostics"]
                report["diagnostics"].extend(d for d in diagnostics if d["severity"] == "error")
                if not any(d["severity"] == "error" for d in report["diagnostics"]):
                    report.update(status="converted", output_path=str(output_path.resolve()),
                                  output_sha256=hashlib.sha256(candidate.read_bytes()).hexdigest())
            except Exception:
                candidate.unlink(missing_ok=True)
                raise
    if rubric_contract_absent and sources.get("rubrics") and not has_refs and proven_absent:
        report["ledger"].append(dict(source="rubrics", disposition="not_applicable",
                                     detail="Absent contract is consistent with no associations and no rubric payload in hashed inventory"))
    report["ledger_path"] = str(ledger_path.resolve())
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(prefix=".wizard-conversion-", suffix=".json", dir=output_path.parent, delete=False) as temporary:
        staged_ledger = Path(temporary.name)
    ledger_published = False
    try:
        staged_ledger.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        if force:
            if candidate is not None and report["status"] == "converted":
                candidate.replace(output_path)
            staged_ledger.replace(ledger_path)
        else:
            # A same-directory hard link publishes complete bytes atomically
            # and refuses a concurrent destination, unlike Path.replace().
            os.link(staged_ledger, ledger_path)
            ledger_published = True
            if candidate is not None and report["status"] == "converted":
                os.link(candidate, output_path)
    except Exception:
        if ledger_published and ledger_path.exists() and os.path.samestat(ledger_path.stat(), staged_ledger.stat()):
            ledger_path.unlink()
        raise
    finally:
        staged_ledger.unlink(missing_ok=True)
        if candidate is not None:
            candidate.unlink(missing_ok=True)
    return report
