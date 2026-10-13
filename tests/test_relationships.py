import hashlib
import zipfile

import pytest
from docx import Document

import synth
from blueprint_packager.docx_ingest import read_docx
from blueprint_packager.relationships import relationship_report, rename_activity
from test_pipeline import model_with_joins, errors


def test_rename_preview_and_apply_preserve_narrative_and_docx_parts(tmp_path):
    old, new = "Week 1 Assignment: Report", "Week 1 Assignment: Investigation"
    source = synth.canonical(tmp_path / "source.docx", overview=["Page: Week 1 Overview", "Page heading: Week 1",
        "Read [link to assignment: Week 1 Assignment: Report | the report].",
        "The narrative refers to Week 1 Assignment: Report and stays as authored."])
    before = source.read_bytes()
    output = tmp_path / "renamed.docx"
    preview = rename_activity(source, output, kind="assignment", old=old, new=new)
    assert preview["status"] == "preview" and not output.exists()
    assert len(preview["changes"]) == 5
    report = relationship_report(model_with_joins(source))
    refs = next(a["references"] for a in report["activities"] if a["title"] == old)
    assert {r["role"] for r in refs} == {"activity declaration", "gradebook declaration", "grading-plan row", "rubric relationship", "internal link"}
    applied = rename_activity(source, output, kind="assignment", old=old, new=new, apply=True,
                              authorized_by="Synthetic operator", reason="Requested exact title update")
    assert applied["status"] == "applied" and source.read_bytes() == before
    assert not errors(model_with_joins(output))
    text = "\n".join(p.text() for p in read_docx(output).paras())
    assert "The narrative refers to Week 1 Assignment: Report and stays as authored." in text
    assert "[link to assignment: Week 1 Assignment: Investigation | the report]" in text
    with zipfile.ZipFile(source) as a, zipfile.ZipFile(output) as b:
        assert a.namelist() == b.namelist()
        for name in a.namelist():
            if name != "word/document.xml":
                assert a.read(name) == b.read(name)


def test_rename_requires_exact_target_authorization_and_new_copy(tmp_path):
    source = synth.canonical(tmp_path / "source.docx")
    options = dict(kind="assignment", old="Week 1 Assignment: Report", new="New exact title")
    with pytest.raises(ValueError, match="requires"):
        rename_activity(source, tmp_path / "new.docx", apply=True, **options)
    with pytest.raises(ValueError, match="different output"):
        rename_activity(source, source, **options)
    with pytest.raises(ValueError, match="found 0"):
        rename_activity(source, tmp_path / "new.docx", **{**options, "old": "Report"})


def test_inherited_word_list_styles_are_semantic_lists(tmp_path):
    doc = Document()
    doc.add_paragraph("One native bullet", style="List Bullet")
    doc.add_paragraph("One native numbered item", style="List Number")
    doc.add_paragraph("An ordinary paragraph")
    path = tmp_path / "lists.docx"
    doc.save(path)
    ingested = read_docx(path)
    paragraphs = list(ingested.paras())
    assert ingested.list_format(paragraphs[0]) == "bullet"
    assert ingested.list_format(paragraphs[1]) == "decimal"
    assert paragraphs[2].list_num_id is None
