import json
from pathlib import Path

import pytest
from docx import Document

import synth
from blueprint_packager.authoring import absolute_web_url, readiness, video_player_url
from blueprint_packager.build import Builder, BuildError
from blueprint_packager.blank import write_blank
from blueprint_packager.html_render import RenderContext, Template, render_blocks
from blueprint_packager.reader import read_model
from test_pipeline import SETTINGS, REPO_ROOT, model_with_joins
from blueprint_packager.verify import _media_check, verify_build


def test_unknown_heading_cannot_hide_content_from_build(tmp_path):
    path = synth.canonical(tmp_path / "source.docx")
    doc = Document(path)
    before = next(p for p in doc.paragraphs if p.text.startswith("Before Week 1"))
    before.insert_paragraph_before("Getting help", "Heading 1")
    before.insert_paragraph_before("This sentence must never disappear.")
    doc.save(path)
    model = model_with_joins(path)
    assert not model["source_accounting"]["complete"]
    assert any(e["text"] == "This sentence must never disappear." and e["disposition"] == "unresolved"
               for e in model["source_accounting"]["entries"])
    with pytest.raises(BuildError, match="error"):
        Builder(model, {**SETTINGS, "build_mode": "preview"}, REPO_ROOT).run(tmp_path / "out", force=False)


def test_missing_overview_is_diagnostic_not_crash(tmp_path):
    path = synth.canonical(tmp_path / "missing.docx", overview=["This overview has no Page declaration."])
    model = model_with_joins(path)
    assert "OVERVIEW_PAGE" in {d["code"] for d in model["diagnostics"]}


def test_trailing_week_row_and_unrecognized_front_table_are_accounted_as_unresolved(tmp_path):
    path = synth.canonical(tmp_path / "extra.docx")
    doc = Document(path)
    week = next(t for t in doc.tables if t.cell(0, 0).text.startswith("Overview:"))
    week.add_row().cells[0].text = "An extra row with authored content."
    doc.save(path)
    model = read_model(path)
    assert any(d["code"] == "WEEK_TABLE_SHAPE" for d in model["diagnostics"])
    assert any(e["disposition"] == "unresolved" and e["text"].startswith("An extra row") for e in model["source_accounting"]["entries"])


def test_guidance_explicitly_excluded_and_ordinary_parentheses_preserved(tmp_path):
    path = synth.canonical(tmp_path / "guidance.docx", overview=["[authoring: Keep this out of the course.]",
        "Page: Week 1 Overview", "Page heading: Week 1", "(This is ordinary authored prose.)", "[placeholder: Write the introduction]"])
    model = model_with_joins(path)
    blocks = model["weeks"][0]["overview_page"]["blocks"]
    assert any(b.get("runs", [{}])[0].get("text") == "(This is ordinary authored prose.)" for b in blocks)
    assert any(e["disposition"] == "authoring_guidance" for e in model["source_accounting"]["entries"])
    report = readiness(model)
    assert [f["code"] for f in report["findings"]] == ["AUTHORING_PLACEHOLDER"]


@pytest.mark.parametrize("url", ["https://", "javascript:alert(1)", "https://example.org/a b", "https://user:pass@example.org/", "https://[invalid"])
def test_incomplete_or_unsafe_web_addresses_rejected(url):
    assert not absolute_web_url(url)


@pytest.mark.parametrize("url,expected", [
    ("https://www.youtube.com/watch?v=abcdefghijk", "https://www.youtube.com/embed/abcdefghijk"),
    ("https://youtu.be/abcdefghijk", "https://www.youtube.com/embed/abcdefghijk"),
    ("https://youtu.be/abcdefghijk?t=1m23s", "https://www.youtube.com/embed/abcdefghijk?start=83"),
    ("https://vimeo.com/123456", "https://player.vimeo.com/video/123456"),
    ("https://player.vimeo.com/video/123456?h=secret", "https://player.vimeo.com/video/123456?h=secret"),
    ("https://www.youtube.com.evil.example/embed/abcdefghijk", None),
    ("https://www.youtube.com/embed/", None),
    ("https://institution.example/player/123", None),
])
def test_player_url_recognition(url, expected):
    assert video_player_url(url) == expected


def test_video_and_placeholder_brackets_round_trip_into_html(tmp_path):
    path = synth.canonical(tmp_path / "video.docx", overview=["Page: Week 1 Overview", "Page heading: Week 1",
        "[video: Weekly introduction | https://youtu.be/abcdefghijk]", "[video placeholder: Demonstration]"])
    model = model_with_joins(path)
    blocks = model["weeks"][0]["overview_page"]["blocks"]
    markup = render_blocks(blocks, RenderContext(Template.load("composite-navy"), "page"))
    assert 'src="https://www.youtube.com/embed/abcdefghijk"' in markup
    assert 'title="Weekly introduction"' in markup
    assert markup.count("<iframe") == 1
    assert "Video to be added: Demonstration" in markup
    assert not readiness(model)["ready"]
    assert readiness(model, {"retain_video_placeholders": ["Demonstration"]})["ready"]


def test_unfinished_blank_build_requires_explicit_preview(tmp_path):
    path = write_blank(tmp_path / "blank.docx", course_number="SYN 101", course_title="Synthetic", term="Sample", weeks=1)
    model = model_with_joins(path)
    with pytest.raises(BuildError, match="draft"):
        Builder(model, SETTINGS, REPO_ROOT).run(tmp_path / "ready", force=False)
    result = Builder(model, {**SETTINGS, "build_mode": "preview"}, REPO_ROOT).run(tmp_path / "preview", force=False)
    assert result["receipt"]["extensions"]["build_mode"] == "preview"
    report = json.loads((tmp_path / "preview/readiness_report.json").read_text(encoding="utf-8"))
    assert report["state"] == "draft"


def test_input_contract_version_and_unsafe_error_override_fail_clearly(tmp_path):
    model = model_with_joins(synth.canonical(tmp_path / "source.docx"))
    with pytest.raises(BuildError, match="schema"):
        Builder({**model, "schema": "coursecraft.blueprint/4"}, SETTINGS, REPO_ROOT)
    with pytest.raises(BuildError, match="allow_model_errors"):
        Builder(model, {**SETTINGS, "allow_model_errors": True}, REPO_ROOT)


def test_media_verification_detects_changed_player_and_duplicate_loss():
    block = {"kind": "video", "title": "A lesson", "href": "https://youtu.be/abcdefghijk", "src": "p001"}
    html = '<iframe src="https://www.youtube.com/embed/abcdefghijk" title="A lesson"></iframe>'
    assert _media_check([block], html, [], "page") == (1, [])
    assert _media_check([block], html.replace("abcdefghijk", "zyxwvutsrqp"), [], "page")[1]
    assert _media_check([block, block], html, [], "page")[1]


def test_video_package_verifies_as_player_and_intentional_placeholder(tmp_path):
    source = synth.canonical(tmp_path / "video.docx", overview=["Page: Week 1 Overview", "Page heading: Week 1",
        "[video: Weekly introduction | https://youtu.be/abcdefghijk]", "[video placeholder: Demonstration]"])
    model = model_with_joins(source)
    settings = {**SETTINGS, "retain_video_placeholders": ["Demonstration"]}
    out = tmp_path / "video-build"
    Builder(model, settings, REPO_ROOT).run(out, force=False)
    assert verify_build(out, model) == 0
    report = json.loads((out / "verification_report.json").read_text(encoding="utf-8"))
    media = next(c for c in report["checks"] if c["check"].startswith("every video"))
    assert media["passed"] and "2 videos" in media["detail"]


def test_exact_legacy_guidance_is_scoped_to_its_authoring_region(tmp_path):
    prompt = "One row per module. The first line of a row names the module; the rest are its pages and links."
    source = synth.canonical(tmp_path / "legacy-guidance.docx", overview=["Page: Week 1 Overview", "Page heading: Week 1", prompt])
    doc = Document(source)
    heading = next(p for p in doc.paragraphs if p.text.startswith("Week 1 -"))
    heading.insert_paragraph_before(prompt)
    doc.save(source)
    model = model_with_joins(source)
    matches = [e for e in model['source_accounting']['entries'] if e['text'] == prompt]
    assert {e['disposition'] for e in matches} == {'content', 'authoring_guidance'}
    assert not [d for d in model['diagnostics'] if d['severity'] == 'error']


def test_legacy_front_matter_prompt_still_marks_draft(tmp_path):
    source = synth.canonical(tmp_path / 'old-prompt.docx')
    doc = Document(source)
    doc.tables[0].cell(1, 0).text = '(Catalog description.)'
    doc.save(source)
    report = readiness(model_with_joins(source))
    assert not report['ready']
    assert any(f['code'] == 'TEMPLATE_SCAFFOLD' and 'Catalog description' in f['message'] for f in report['findings'])
