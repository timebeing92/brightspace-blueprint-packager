"""Generate a reviewable authoring blueprint with explicit guidance and placeholders.

The minimal form contains writing areas and selected activity scaffolds. The worked
example adds clearly named sample rubrics and sample content. Every structural cue
is text-based so it survives a Word / Google Docs round trip.
"""

from __future__ import annotations

import io
import os
import re
import shutil
import tempfile
from datetime import datetime, timezone
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from . import vocabulary as V
from .authoring import AUTHOR_NOTE_PREFIX, WRITING_PROMPTS

NAVY = RGBColor(0x10, 0x39, 0x45)
BRAND = RGBColor(0x17, 0x4B, 0x5C)
GREY = RGBColor(0x5B, 0x6B, 0x72)
LINK = RGBColor(0x07, 0x5F, 0x7A)
LABEL_FILL = "EDF4F6"

# Panels are shaded the way Blueprint Wizard shades them in an export blueprint, so the
# two documents look alike: dropdowns lavender, callouts blue; the gold callout takes the
# composite template's warm aside colours. The keyword on each bracket line is coloured to match.
PANELS = {
    "dropdown": {"fill": "F4F2FA", "ink": RGBColor(0x5B, 0x3A, 0x8C)},
    "callout": {"fill": "F2F6FA", "ink": NAVY},
    "callout gold": {"fill": "FBF1E9", "ink": RGBColor(0x8B, 0x4A, 0x27)},
    "callout warm": {"fill": "FBF1E9", "ink": RGBColor(0x8B, 0x4A, 0x27)},
    "timeline": {"fill": "F2F6FA", "ink": NAVY},
    "box": {"fill": "F2F6FA", "ink": NAVY},
}
_LINK_RE = re.compile(r"\[link to [^\]]+\]")
_OPENER_RE = re.compile(r"^\[\s*([a-z]+)((?:\s+[a-z_]+=[^\s\]:]+)*)")

# Guidance is explicitly excluded by the reader; placeholders remain visible and
# are reported by readiness checks. Both forms survive ordinary document conversion.
GUIDE = [
    ("Start with one week", [
        ("Draft the learning experience.", "Confirm the course outcomes and assessment plan with your instructional designer. Then complete one week's overview, objectives, resources, and student instructions before repeating the pattern."),
        ("Replace a writing prompt.", "Select and replace the entire prompt sentence. Type or paste your text in its place, then press Enter to continue with more paragraphs. The writing area grows with your content."),
        ("Keep the structure in place.", "Keep shaded section labels and field names such as Page: and Points:. Change the values after the colon with your designer. Use Heading 2 or deeper for headings within student content."),
        ("Return your draft for review.", "Save or download as Word (.docx) and send it to your instructional designer. Review the learner preview together after checking activity names, points, rubrics, links, and dates."),
    ]),
]

SECTION_TITLES = {
    "overview": "Overview:",
    "objectives": "Learning Objectives:",
    "resources": "Assigned Reading and Multimedia:",
    "assignments": "Assignment(s) and Instructions:",
    "discussions": "Discussion Board Prompts:",
    "checklist": "Checklist",
}


# ------------------------------------------------------------------ styling helpers
def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _shade_paragraph(para, fill: str) -> None:
    p_pr = para._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def _panel_for(line: str) -> str | None:
    """'dropdown', 'callout', 'callout gold', … for an opening bracket line; None otherwise."""
    match = _OPENER_RE.match(line.strip())
    if not match or line.startswith("[/"):
        return None
    keyword, attrs = match.group(1), match.group(2) or ""
    tone = re.search(r"tone=([a-z]+)", attrs)
    key = f"{keyword} {tone.group(1)}" if tone else keyword
    return key if key in PANELS else (keyword if keyword in PANELS else None)


def _write_text_with_links(para, text: str, colour=None) -> None:
    """Plain text, with every [link to …] bracket in link blue and underlined."""
    pos = 0
    for match in _LINK_RE.finditer(text):
        if match.start() > pos:
            run = para.add_run(text[pos:match.start()])
            if colour is not None:
                run.font.color.rgb = colour
        link = para.add_run(match.group(0))
        link.font.color.rgb = LINK
        link.underline = True
        pos = match.end()
    if pos < len(text):
        run = para.add_run(text[pos:])
        if colour is not None:
            run.font.color.rgb = colour


def _cell_margins(table, top: int = 80, bottom: int = 80, left: int = 140, right: int = 140) -> None:
    """Breathing room inside every cell, in twentieths of a point."""
    tbl_pr = table._tbl.tblPr
    margins = OxmlElement("w:tblCellMar")
    for side, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tbl_pr.append(margins)


def _borders(table, color: str = "C7D8DC") -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)
        borders.append(node)
    tbl_pr.append(borders)


def _is_prompt(line: str) -> bool:
    return line.startswith("(") and line.endswith(")")


def _is_structure(line: str) -> bool:
    return V.parse_directive(line) is not None or V.parse_bracket_block(line) is not None


def _write_line(para, line: str, panel: str | None = None) -> None:
    """One line into a paragraph.

    Prompts are grey italic; structure keywords bold navy; bracket lines take their panel's
    ink; a line inside a panel is shaded with the panel's fill; [link to …] is link blue."""
    para.paragraph_format.space_after = Pt(3)
    if panel:
        _shade_paragraph(para, PANELS[panel]["fill"])
    if line.startswith(("[authoring:", AUTHOR_NOTE_PREFIX)):
        if line.startswith(AUTHOR_NOTE_PREFIX):
            lead = para.add_run(AUTHOR_NOTE_PREFIX)
            lead.bold = True
            run = para.add_run(line[len(AUTHOR_NOTE_PREFIX):])
            lead.font.size = Pt(9)
            lead.font.color.rgb = GREY
        else:
            run = para.add_run(line)
        run.font.size = Pt(9)
        run.font.color.rgb = GREY
        return
    if line.startswith("[placeholder:") or line in WRITING_PROMPTS.values():
        run = para.add_run(line)
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.italic = False
        run.bold = False
        para.paragraph_format.space_before = Pt(5)
        para.paragraph_format.space_after = Pt(8)
        return
    if line.startswith("["):
        run = para.add_run(line)
        key = _panel_for(line) or panel or "callout"
        run.font.color.rgb = PANELS[key]["ink"]
        run.bold = True
        return
    if _is_prompt(line):
        run = para.add_run(line)
        run.italic = True
        run.font.color.rgb = GREY
        return
    if _is_structure(line) and ":" in line:
        if line.startswith(("Assignment:", "Discussion:")):
            para.paragraph_format.space_before = Pt(4)
            para.paragraph_format.keep_with_next = True
        keyword, rest = line.split(":", 1)
        lead = para.add_run(keyword + ":")
        lead.bold = True
        lead.font.color.rgb = NAVY
        _write_text_with_links(para, rest)
        return
    if line.startswith("(An internal link reads like this:)"):
        note, rest = line.split(")", 1)
        prompt = para.add_run(note + ")")
        prompt.italic = True
        prompt.font.color.rgb = GREY
        _write_text_with_links(para, rest)
        return
    _write_text_with_links(para, line)


def _fill(cell, lines: list[str], label: bool = False) -> None:
    first = True
    activities = 0
    panel: str | None = None  # the panel we are inside, so its lines take its fill
    for line in lines:
        para = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        if label:
            run = para.add_run(line)
            run.bold = True
            run.font.color.rgb = NAVY
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.space_before = Pt(2)
            continue
        opened = _panel_for(line)
        if opened:
            panel = opened
        if (line.startswith("[placeholder: Learning objective") or line.startswith("[placeholder: Course learning outcome")
                or line in (WRITING_PROMPTS["learning_objective"], WRITING_PROMPTS["course_learning_outcome"])):
            para.style = "List Bullet"
            # Direct numbering survives readers that do not resolve style inheritance.
            num_pr = para.style.element.xpath("./w:pPr/w:numPr")
            if num_pr:
                para._p.get_or_add_pPr().append(deepcopy(num_pr[0]))
        _write_line(para, line, panel)
        if line.startswith(("Assignment:", "Discussion:")):
            if activities:
                para.paragraph_format.space_before = Pt(14)
            activities += 1
        if line.startswith("[/"):
            panel = None



def _no_split(row) -> None:
    """A row never breaks across pages, so a shaded label is not left alone at a page's top."""
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def _one_column(doc, rows: list[list[str]], label_rows: set[int] | None = None, *, spacer: bool = True):
    table = doc.add_table(rows=len(rows), cols=1)
    table.autofit = True
    _borders(table)
    _cell_margins(table)
    for index, lines in enumerate(rows):
        is_label = index in (label_rows or set())
        row = table.rows[index]
        cell = row.cells[0]
        if is_label:
            _no_split(row)
            _shade(cell, LABEL_FILL)
        _fill(cell, lines, label=is_label)
        if is_label:
            for para in cell.paragraphs:
                para.paragraph_format.keep_with_next = True
    if spacer:
        gap = doc.add_paragraph()
        gap.paragraph_format.space_after = Pt(4)
        gap.paragraph_format.space_before = Pt(0)
        gap.paragraph_format.line_spacing = Pt(1)
        gap.add_run().font.size = Pt(1)
    return table


def _style_document(doc) -> None:
    section = doc.sections[0]
    section.left_margin = section.right_margin = Inches(1)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    defaults = doc.styles.element.xpath("./w:docDefaults/w:rPrDefault/w:rPr/w:rFonts")
    for fonts in defaults:
        fonts.attrib.clear()
        for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{slot}"), "Arial")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.1
    for name, size, before, after in (("Title", 26, 0, 6), ("Heading 1", 18, 20, 8), ("Heading 2", 14, 12, 4), ("Heading 3", 12, 10, 3)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = NAVY
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    # Remove the Word default Title border, including inherited direct style residue.
    for name in ("Title", "Subtitle"):
        for border in doc.styles[name].element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)


def _guide(doc) -> None:
    doc.add_paragraph("Use this document to draft the content students will read and to agree the course setup with your instructional designer.")
    for heading, lines in GUIDE:
        doc.add_paragraph(heading, style="Heading 2")
        for lead, text in lines:
            item = doc.add_paragraph(style="List Number")
            item.paragraph_format.space_after = Pt(7)
            item.add_run(lead + " ").bold = True
            item.add_run(text)
    doc.add_paragraph("What writing looks like", style="Heading 2")
    p = doc.add_paragraph()
    p.add_run("Prompt to replace   ").bold = True
    p.add_run(WRITING_PROMPTS["week_introduction"])
    p = doc.add_paragraph()
    p.add_run("Your replacement   ").bold = True
    p.add_run("This week, you will compare two ways to evaluate a source.")
    doc.add_paragraph("Continue with as many paragraphs, lists, links, and headings as needed. Your prose is ordinary document text; brackets are only used for optional features such as videos and expandable sections.")
    doc.add_paragraph("A few useful distinctions", style="Heading 2")
    for lead, text in (
        ("Author notes", "help you complete the document and stay out of the course."),
        ("Course reference fields", "record the description, materials and course outcomes; they do not automatically create student pages."),
        ("Authoring reference", "at the end contains separate examples to copy when you need another activity, a video, a panel, or a link."),
    ):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(lead + " ").bold = True
        p.add_run(text)

    doc.add_paragraph("Choose the features that help you", style="Heading 2")
    doc.add_paragraph("Designers and SMEs are welcome to use the optional design tokens (dropdowns, callouts, and videos) and shortcuts that connect course objects such as assignments, grade items, and rubrics. If you are more comfortable, use ordinary text and defer styling or connections for manual completion in Brightspace. Record deferred work in Extraction Notes; your designer can remove any unused connection settings.")


def _reference_example(doc, heading: str, purpose: str, lines: list[str], after: str = "") -> None:
    doc.add_paragraph(heading, style="Heading 3")
    doc.add_paragraph(purpose)
    label = doc.add_paragraph("Copy the following lines")
    label.paragraph_format.keep_with_next = True
    label.runs[0].bold = True
    for index, line in enumerate(lines):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        _write_line(p, line)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = index < len(lines) - 1
    if after:
        doc.add_paragraph(after)


def _reference(doc) -> None:
    heading = doc.add_paragraph(V.H_AUTHORING_REFERENCE, style="Heading 1")
    heading.paragraph_format.page_break_before = True
    doc.add_paragraph("This section is for authors and stays out of the course. Copy only the example lines into the relevant writing area. Replace the example titles and writing prompts; keep the field names and functional brackets.")
    doc.add_paragraph("Add or change activities", style="Heading 2")
    _reference_example(doc, "Add an assignment", "Paste into the week's Assignment(s) and Instructions cell, after the preceding activity.", [
        "Assignment: Week 2 Assignment: Source comparison", "Points: 20", "Gradebook item: Week 2 Assignment: Source comparison",
        WRITING_PROMPTS["assignment_instructions"]],
        "Designer: add a grading-plan row with that exact name and 20 points. If using a rubric, add Rubric: followed by its exact name above the instructions, and add the activity to that rubric's Used by list.")
    _reference_example(doc, "Add a discussion", "Paste into the week's Discussion Board Prompts cell. The next Discussion: line begins a separate activity.", [
        "Discussion: Week 2 Discussion: Comparing evidence", "Points: 10", "Gradebook item: Week 2 Discussion: Comparing evidence",
        "Forum: Discussions", WRITING_PROMPTS["discussion_prompt"]],
        "Designer: add a matching 10-point grading-plan row. For an ungraded activity, omit Points: and Gradebook item: and do not add a grading-plan row.")
    doc.add_paragraph("Rename or remove an activity", style="Heading 3")
    doc.add_paragraph("Copy or delete from the Assignment: or Discussion: line through the final instruction paragraph, stopping before the next activity. Leave the shaded weekly section label in place. Have your designer update the grading row, rubric references, checklist wording and links. Use the Packager's rename preview when changing an established title.")
    doc.add_paragraph("Prefer to finish setup in Brightspace?", style="Heading 3")
    doc.add_paragraph("Design tokens and connection shortcuts are optional. You can style content and connect course objects by hand in Brightspace after import. Record that work in Extraction Notes. Your designer should remove unused connection settings and matching references before packaging; any connection left in the blueprint must still match a real target exactly. See the Authoring Guide for the grading and rubric details.")

    heading = doc.add_paragraph("Expandable sections and callouts", style="Heading 2")
    heading.paragraph_format.page_break_before = True
    doc.add_paragraph("Optional page elements. Put each bracket line in its own paragraph. Write ordinary content between the opening and closing lines; you can use multiple paragraphs, lists and links.")
    _reference_example(doc, "Expandable section", "Students open this section to reveal the text. Change Additional detail to the title they should see.", [
        "[dropdown: Additional detail]", WRITING_PROMPTS["panel_content"], "[/dropdown]"])
    _reference_example(doc, "Blue callout", "Use a short callout for a note or reminder that supports the surrounding content.", [
        "[callout]", WRITING_PROMPTS["panel_content"], "[/callout]"])
    _reference_example(doc, "Warm callout", "The same kind of note, with the alternative warm colour treatment.", [
        "[callout tone=gold]", WRITING_PROMPTS["panel_content"], "[/callout]"])

    heading = doc.add_paragraph("Videos and links", style="Heading 2")
    heading.paragraph_format.page_break_before = True
    doc.add_paragraph("Use these forms only where you want the feature to appear. Each video bracket occupies one paragraph; an internal-link bracket can sit inside a sentence.")
    _reference_example(doc, "Video with a known URL", "Replace Weekly introduction with an accessible title and replace VIDEO_ID with the video's YouTube ID, or paste a supported Vimeo URL.", [
        "[video: Weekly introduction | https://www.youtube.com/watch?v=VIDEO_ID]"])
    _reference_example(doc, "Video to add later", "Reserve a visible place for a video when its URL is not yet available.", [
        "[video placeholder: Weekly introduction]"], "Before packaging, add the URL or have your designer explicitly retain the named video placeholder.")
    _reference_example(doc, "Link to an activity", "Replace the target title with the activity's exact title. Words after the vertical bar are the link text students see.", [
        "[link to assignment: Week 2 Assignment: Source comparison | Open the assignment]"],
        "The same form accepts discussion, checklist, page, or quiz when that target exists.")
    doc.add_paragraph("Before the designer builds", style="Heading 3")
    doc.add_paragraph("Review outcome alignment, student workload, task and submission clarity, accessibility, materials access, and assessment criteria. Confirm points deliberately, including any intended zero-point or ungraded work. The generated checklist covers activities; add reading, viewing or practice tasks when needed. Technical readiness alone does not establish instructional completeness or approval.")


def _repeat_header(row) -> None:
    _no_split(row)
    header = OxmlElement("w:tblHeader")
    row._tr.get_or_add_trPr().append(header)


def _table_widths(table, widths: tuple[float, ...]) -> None:
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)


def _activity_counts(weeks: int, assignments: int, discussions: int,
                     overrides: dict[int, dict[str, int]] | None) -> dict[int, dict[str, int]]:
    def count(value, label):
        if isinstance(value, bool) or not isinstance(value, int) or value < 0:
            raise ValueError(f"{label} must be a non-negative integer")
        return value

    if isinstance(weeks, bool) or not isinstance(weeks, int) or not 1 <= weeks <= 52:
        raise ValueError("weeks must be an integer from 1 to 52")
    defaults = {"assignments": count(assignments, "assignments_per_week"),
                "discussions": count(discussions, "discussions_per_week")}
    result = {n: dict(defaults) for n in range(1, weeks + 1)}
    for n, values in (overrides or {}).items():
        if isinstance(n, bool) or not isinstance(n, int) or n not in result:
            raise ValueError(f"week_counts key {n!r} must identify a week from 1 to {weeks}")
        if not isinstance(values, dict) or set(values) - set(defaults):
            raise ValueError(f"week_counts[{n}] accepts only assignments and discussions")
        for key, value in values.items():
            result[n][key] = count(value, f"week_counts[{n}][{key}]")
    return result


def _activity_title(n: int, kind: str, index: int, count: int) -> str:
    sequence = f" {index}" if count > 1 else ""
    return f"Week {n} {kind}{sequence}: (title)"


def _activity_lines(n: int, kind: str, titles: list[str], worked_example: bool) -> list[str]:
    points = (10 if kind == "Assignment" else 5) if worked_example else 0
    lines = [] if titles else [f"Author note: Optional. Add a {kind}: block here if needed; see Authoring reference."]
    for title in titles:
        lines += [f"{kind}: {title}", f"Points: {points} pts", f"Gradebook item: {title}"]
        if worked_example:
            lines.append(f"Rubric: Sample {kind} Rubric")
        if kind == "Discussion":
            lines.append("Forum: Discussions")
        lines.append("Author note: Explain the task, what to submit, and how success will be judged." if kind == "Assignment" else
                     "Author note: State the prompt and expectations for the initial post and replies.")
        lines.append(WRITING_PROMPTS["assignment_instructions" if kind == "Assignment" else "discussion_prompt"])
    return lines


def _save_document(doc, path: Path, *, force: bool) -> None:
    """Write safely; an explicit overwrite retains a unique copy of the old file."""
    data = io.BytesIO()
    doc.save(data)
    if not force:
        # Exclusive creation also protects against a file appearing after the initial check.
        with path.open("xb") as target:
            target.write(data.getvalue())
        return
    if path.exists():
        stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
        backup = path.with_name(f"{path.stem}.backup-{stamp}{path.suffix}")
        with path.open("rb") as source, backup.open("xb") as target:
            shutil.copyfileobj(source, target)
        shutil.copystat(path, backup)
    fd, temp_name = tempfile.mkstemp(prefix=f".{path.stem}-", suffix=path.suffix, dir=path.parent)
    try:
        with os.fdopen(fd, "wb") as target:
            target.write(data.getvalue())
        os.replace(temp_name, path)
    finally:
        Path(temp_name).unlink(missing_ok=True)


def write_blank(path: Path, *, course_number: str, course_title: str, term: str, weeks: int = 8,
                topics: list[str] | None = None, assignments_per_week: int = 1,
                discussions_per_week: int = 1, week_counts: dict[int, dict[str, int]] | None = None,
                worked_example: bool = False, force: bool = False) -> Path:
    """Create a new authoring blueprint; refuse replacement unless ``force`` is explicit.

    ``week_counts={2: {"assignments": 3, "discussions": 0}}`` overrides common
    defaults for that week. The minimal form uses zero-point placeholders and
    no rubric assumptions. ``worked_example`` adds sample 10/5-point activities,
    two example rubrics, and demonstrative content blocks in the first week.
    """
    path = Path(path)
    if path.exists() and not force:
        raise FileExistsError(f"Blueprint already exists: {path}. Choose another filename or use force to retain a backup and replace it.")
    counts = _activity_counts(weeks, assignments_per_week, discussions_per_week, week_counts)
    topics = list(topics or [])
    titles = {n: {kind: [_activity_title(n, kind, i, counts[n][key])
                         for i in range(1, counts[n][key] + 1)]
                  for key, kind in (("assignments", "Assignment"), ("discussions", "Discussion"))}
              for n in counts}
    doc = Document()
    _style_document(doc)
    footer = doc.sections[0].footer.paragraphs[0]
    footer.paragraph_format.space_after = Pt(0)
    run = footer.add_run(f"{course_number}  |  {term}  |  Page ")
    run.font.size = Pt(8)
    run.font.color.rgb = GREY
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.add_paragraph("Course blueprint", style="Title")
    sub = doc.add_paragraph(f"{course_number} {course_title} – {term} Blueprint")
    sub.runs[0].font.size = Pt(13)
    sub.runs[0].font.color.rgb = GREY
    mode = "Worked example with sample activities and rubrics" if worked_example else "Blank for collaborative course authoring"
    _write_line(doc.add_paragraph(), mode)
    _guide(doc)
    doc.add_paragraph().paragraph_format.page_break_before = True

    doc.add_paragraph("Course reference and assessment plan", style="Heading 2")
    _write_line(doc.add_paragraph(), "Author note: These course reference fields stay in the blueprint. To show this material to students, your designer must also place it in an explicit course page. Delete a writing prompt when an optional field is not needed.")
    _one_column(doc, [[V.FRONT_LABELS["course_description"] + " (catalog)"],
                     [WRITING_PROMPTS["course_description"]]], {0})
    _one_column(doc, [[V.FRONT_LABELS["required_materials"]],
                     [WRITING_PROMPTS["required_materials"]]], {0})
    _one_column(doc, [[V.FRONT_LABELS["course_learning_outcomes"]],
                     ["Author note: Use one course outcome per bullet.",
                      WRITING_PROMPTS["course_learning_outcome"]]], {0})
    doc.add_paragraph("Grading plan for designer review", style="Heading 2")
    _write_line(doc.add_paragraph(), "Author note: The Assignment column lists all graded activities, including discussions. Set points deliberately; the starter 0 values are not a grading recommendation. Keep names and points identical in this table and the activity settings.")
    grade_rows = [(title, (10 if kind == "Assignment" else 5) if worked_example else 0)
                  for n in titles for kind in ("Assignment", "Discussion") for title in titles[n][kind]]
    table = doc.add_table(rows=len(grade_rows) + 2, cols=2)
    _borders(table)
    _cell_margins(table, top=60, bottom=60)
    _table_widths(table, (5.25, 1.25))
    _repeat_header(table.rows[0])
    for cell, text in zip(table.rows[0].cells, V.GRADE_PLAN_HEADER):
        _shade(cell, LABEL_FILL)
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.color.rgb = NAVY
    for index, (name, points) in enumerate(grade_rows, start=1):
        table.rows[index].cells[0].paragraphs[0].add_run(name)
        table.rows[index].cells[1].paragraphs[0].add_run(str(points))
        _no_split(table.rows[index])
    total = table.rows[-1]
    for cell in total.cells:
        _shade(cell, LABEL_FILL)
    total.cells[0].paragraphs[0].add_run(V.GRADE_PLAN_TOTAL).bold = True
    total.cells[1].paragraphs[0].add_run(str(sum(p for _n, p in grade_rows))).bold = True
    _no_split(total)

    doc.add_paragraph(V.H_COURSE_INTRO, style="Heading 1")
    for line in ("Page: Welcome", f"Page heading: Welcome to {course_number}", "Placement: Getting Started",
                 WRITING_PROMPTS["welcome"]):
        _write_line(doc.add_paragraph(), line)
    doc.add_paragraph(V.H_BEFORE_WEEK_1, style="Heading 1")
    _one_column(doc, [["Getting Started"],
        [f"{V.D_MODULE}: Getting Started", "Page: Student Resources", "Page heading: Student and Course Resources",
         WRITING_PROMPTS["student_resources"],
         "Author note: To add a syllabus link, insert Link: Syllabus, then URL: followed by its full address on the next line."]], {0})

    for n in range(1, weeks + 1):
        topic = topics[n - 1] if n - 1 < len(topics) else "(topic)"
        week_heading = doc.add_paragraph(f"Week {n} - {topic}", style="Heading 1")
        week_heading.paragraph_format.page_break_before = True
        week_heading.paragraph_format.space_before = Pt(0)
        overview = [f"Page: Week {n} Overview", f"Page heading: Week {n} – {topic}",
                    "Author note: Introduce the week's purpose and connect its ideas and activities.",
                    WRITING_PROMPTS["week_introduction"]]
        if worked_example and n == 1:
            overview += ["Author note: These sample panels become course content; adapt or delete each complete block.",
                         "[dropdown: Sample optional detail]", WRITING_PROMPTS["panel_content"], "[/dropdown]",
                         "[callout]", WRITING_PROMPTS["panel_content"], "[/callout]",
                         "[callout tone=gold]", WRITING_PROMPTS["panel_content"], "[/callout]",
                         "[video placeholder: Week 1 introduction]"]
        sources = {
            "overview": overview,
            "objectives": [V.OBJECTIVES_SUBLABEL,
                           "Author note: Use one observable objective per bullet. Press Enter to add another.",
                           WRITING_PROMPTS["learning_objective"]],
            "resources": ["Author note: Identify required or optional resources and their purpose. Delete the prompt if unused.",
                          WRITING_PROMPTS["readings_multimedia"]],
            "assignments": _activity_lines(n, "Assignment", titles[n]["Assignment"], worked_example),
            "discussions": _activity_lines(n, "Discussion", titles[n]["Discussion"], worked_example),
            "checklist": [f"Checklist: Week {n} Checklist", f"{V.D_ITEMS}: {V.ITEMS_FROM_ACTIVITIES}",
                          "Author note: This lists activities automatically. Add reading or practice tasks as extra paragraphs if needed."],
        }
        rows = []
        for key, _prefix, _label in V.SECTION_LABELS:
            rows.extend([[SECTION_TITLES[key]], sources[key]])
        _one_column(doc, rows, set(range(0, len(rows), 2)), spacer=False)

    doc.add_paragraph(V.H_COURSE_ACTIVITIES, style="Heading 1")
    _write_line(doc.add_paragraph(), "Author note: Optional. Add complete activity blocks here only for activities that belong to no week.")
    doc.add_paragraph(V.H_RUBRICS, style="Heading 1")
    _write_line(doc.add_paragraph(), "Author note: Designer setup. Leave this section empty if no rubrics are needed. Otherwise use one Heading 2 per rubric, followed by its Used by: lines and rubric table. Level headers carry percentages; criterion weights total 100%.")
    if worked_example:
        _write_line(doc.add_paragraph(), "Author note: These two-criterion rubrics are samples. Replace their criteria, levels, weights and descriptors with the agreed assessment criteria.")
        for kind in ("Assignment", "Discussion"):
            users = [title for n in titles for title in titles[n][kind]]
            if not users:
                continue
            doc.add_paragraph(f"Sample {kind} Rubric", style="Heading 2")
            for title in users:
                _write_line(doc.add_paragraph(), f"Used by: {kind}: {title}")
            grid = doc.add_table(rows=3, cols=4)
            _borders(grid)
            _cell_margins(grid, top=60, bottom=60)
            _table_widths(grid, (1.75, 1.58, 1.58, 1.59))
            _repeat_header(grid.rows[0])
            for cell, text in zip(grid.rows[0].cells, ("Criterion", "Exemplary — 100%", "Proficient — 80%", "Developing — 50%")):
                _shade(cell, LABEL_FILL)
                run = cell.paragraphs[0].add_run(text)
                run.bold = True
                run.font.color.rgb = NAVY
            for row, first in zip(grid.rows[1:], ("(first criterion) (50%)", "(second criterion) (50%)")):
                for cell, text in zip(row.cells, (first, *(WRITING_PROMPTS["rubric_descriptor"],) * 3)):
                    _write_line(cell.paragraphs[0], text)
    doc.add_paragraph(V.H_NOTES, style="Heading 1")
    _write_line(doc.add_paragraph(), "Author note: Optional notes for your instructional designer. This section stays out of the course.")
    _reference(doc)
    _save_document(doc, path, force=force)
    return path
