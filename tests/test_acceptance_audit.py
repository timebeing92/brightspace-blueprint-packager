"""Independent acceptance regressions for content boundaries and exact renames."""
from __future__ import annotations

import pytest
from docx import Document

import synth
from blueprint_packager.build import Builder, BuildError
from blueprint_packager.docx_ingest import read_docx
from blueprint_packager.relationships import rename_activity
from test_pipeline import SETTINGS, REPO_ROOT, model_with_joins


@pytest.mark.parametrize("surface", ["week", "preweek", "front"])
@pytest.mark.parametrize("placement", ["extra_paragraph", "same_paragraph_suffix"])
def test_unrecognized_label_content_is_not_silently_discarded(tmp_path, surface, placement):
    source = synth.canonical(tmp_path / "source.docx")
    doc = Document(source)
    prefixes = {"week": "Overview:", "preweek": "Getting Started", "front": "COURSE DESCRIPTION"}
    table = next(table for table in doc.tables if table.cell(0, 0).text.startswith(prefixes[surface]))
    cell = table.cell(0, 0)
    sentence = "Unique synthetic authored sentence that must survive."
    if placement == "extra_paragraph":
        cell.add_paragraph(sentence)
    else:
        cell.paragraphs[0].add_run(" " + sentence)
    doc.save(source)
    model = model_with_joins(source)
    assert any(d["severity"] == "error" for d in model["diagnostics"])
    assert any(sentence in entry["text"] and entry["disposition"] == "unresolved"
               for entry in model["source_accounting"]["entries"])
    with pytest.raises(BuildError):
        Builder(model, SETTINGS, REPO_ROOT).run(tmp_path / "out", force=False)


def test_rename_does_not_change_directive_like_narrative(tmp_path):
    old, new = "Week 1 Assignment: Report", "Week 1 Assignment: Investigation"
    body = ["Body introduction.", "[dropdown: A quoted structure example]", f"Assignment: {old}",
            f"Gradebook item: {old}", f"Used by: Assignment: {old}", "[/dropdown]", "Body conclusion."]
    source = synth.canonical(tmp_path / "source.docx", assignment_body=body)
    model = model_with_joins(source)
    assert not [d for d in model["diagnostics"] if d["severity"] == "error"]
    content_ids = {entry["id"] for entry in model["source_accounting"]["entries"] if entry["disposition"] == "content"}
    result = rename_activity(source, tmp_path / "renamed.docx", kind="assignment", old=old, new=new,
                             apply=True, authorized_by="Synthetic operator", reason="Requested exact object rename")
    assert not content_ids.intersection(change["paragraph_id"] for change in result["changes"])
    text = "\n".join(para.text() for para in read_docx(tmp_path / "renamed.docx").paras())
    for line in body:
        assert line in text
    assert not [d for d in result["diagnostics"] if d["severity"] == "error"]


def test_missing_source_id_does_not_bypass_pending_content_readiness(tmp_path):
    from blueprint_packager.authoring import readiness

    source = synth.canonical(tmp_path / "source.docx", overview=[
        "Page: Week 1 Overview", "Page heading: Week 1", "[placeholder: Complete this synthetic draft sentence]"])
    model = model_with_joins(source)
    model["weeks"][0]["overview_page"]["blocks"][0].pop("src")
    assert not readiness(model)["ready"]
    assert any(f["code"] == "AUTHORING_PLACEHOLDER" for f in readiness(model)["findings"])
    with pytest.raises(BuildError, match="draft"):
        Builder(model, SETTINGS, REPO_ROOT).run(tmp_path / "out", force=False)


def test_missing_source_id_does_not_skip_video_readiness_or_verification(tmp_path):
    from blueprint_packager.authoring import readiness
    from blueprint_packager.verify import _media_check

    source = synth.canonical(tmp_path / "source.docx", overview=[
        "Page: Week 1 Overview", "Page heading: Week 1", "[video placeholder: Synthetic pending video]"])
    model = model_with_joins(source)
    block = model["weeks"][0]["overview_page"]["blocks"][0]
    block.pop("src")
    assert not readiness(model)["ready"]
    assert readiness(model, {"retain_video_placeholders": ["Synthetic pending video"]})["ready"]
    count, issues = _media_check([block], "", [], "Overview")
    assert count == 1 and issues


@pytest.mark.parametrize("mutation", ["null_grade_plan", "null_overview", "empty_overview"])
def test_missing_mandatory_shapes_raise_build_errors(tmp_path, mutation):
    source = synth.canonical(tmp_path / "source.docx")
    model = model_with_joins(source)
    if mutation == "null_grade_plan":
        model["grade_plan"] = None
    else:
        model["weeks"][0]["overview_page"] = None if mutation == "null_overview" else {}
    # Contract failure must be an actionable BuildError, not a later TypeError
    # or KeyError escaping from plan().
    with pytest.raises(BuildError):
        Builder(model, SETTINGS, REPO_ROOT).plan()


def test_rename_internal_link_titles_are_case_sensitive(tmp_path):
    old, new = "Week 1 Assignment: Report", "Week 1 Assignment: Investigation"
    invalid = "Read [link to assignment: week 1 assignment: report | a misspelled target]."
    valid = f"Read [LINK TO ASSIGNMENT: {old} | the actual target]."
    source = synth.canonical(tmp_path / "source.docx", overview=[
        "Page: Week 1 Overview", "Page heading: Week 1", invalid, valid])
    preview = rename_activity(source, tmp_path / "new.docx", kind="assignment", old=old, new=new)
    assert not any(change["before"] == invalid for change in preview["changes"])
    assert any(change["before"] == valid and f"[LINK TO ASSIGNMENT: {new}" in change["after"]
               for change in preview["changes"])


def test_rename_collision_cleanup_preserves_other_writers_file(tmp_path, monkeypatch):
    import zipfile
    import blueprint_packager.relationships as relationships

    source = synth.canonical(tmp_path / "source.docx")
    output = tmp_path / "new.docx"
    original_zip = zipfile.ZipFile
    other_writers_bytes = b"Another writer's authored file must remain intact."

    def raced_open(path, *args, **kwargs):
        if path == output and args and args[0] == "x":
            output.write_bytes(other_writers_bytes)
        return original_zip(path, *args, **kwargs)

    monkeypatch.setattr(relationships.zipfile, "ZipFile", raced_open)
    with pytest.raises(FileExistsError):
        rename_activity(source, output, kind="assignment", old="Week 1 Assignment: Report", new="New synthetic title",
                        apply=True, authorized_by="Synthetic operator", reason="Requested exact object rename")
    assert output.read_bytes() == other_writers_bytes
