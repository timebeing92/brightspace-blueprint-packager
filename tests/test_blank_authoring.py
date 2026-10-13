"""Synthetic authoring-template behavior, including safe destinations and activity counts."""
from pathlib import Path

import pytest
from docx import Document

from blueprint_packager.blank import write_blank
from blueprint_packager.authoring import WRITING_PROMPTS
from blueprint_packager.joins import resolve
from blueprint_packager.reader import read_model


BASE = dict(course_number="SYN 101", course_title="Synthetic authoring", term="Test term")


def generate(tmp_path: Path, **kwargs):
    return write_blank(tmp_path / "blueprint.docx", **BASE, **kwargs)


def test_existing_authored_file_is_preserved_unless_force_is_explicit(tmp_path):
    path = generate(tmp_path, weeks=1)
    doc = Document(path)
    doc.add_paragraph("Synthetic authored draft must survive.")
    doc.save(path)
    authored_bytes = path.read_bytes()
    with pytest.raises(FileExistsError, match="already exists"):
        generate(tmp_path, weeks=2)
    assert path.read_bytes() == authored_bytes
    generate(tmp_path, weeks=2, force=True)
    backups = list(tmp_path.glob("blueprint.backup-*.docx"))
    assert len(backups) == 1
    assert backups[0].read_bytes() == authored_bytes
    assert path.read_bytes() != authored_bytes
    generate(tmp_path, weeks=3, force=True)
    assert len(list(tmp_path.glob("blueprint.backup-*.docx"))) == 2


def test_minimal_blank_has_no_rubric_assumptions_or_sample_panels(tmp_path):
    model = read_model(generate(tmp_path, weeks=1))
    assert model["rubrics"] == []
    assert len(model["weeks"][0]["assignments"]) == 1
    assert model["weeks"][0]["assignments"][0]["points"] == 0
    assert not any(b["kind"] in {"dropdown", "visual"} for b in model["weeks"][0]["overview_page"]["blocks"])
    assert not model["links"]  # no unfinished URL is generated


def test_per_week_activity_counts_generate_matching_grade_and_rubric_references(tmp_path):
    model = read_model(generate(tmp_path, weeks=3, assignments_per_week=2,
                               discussions_per_week=1, week_counts={2: {"assignments": 0, "discussions": 3},
                                                                    3: {"discussions": 0}},
                               worked_example=True))
    assert [(len(w["assignments"]), len(w["discussions"])) for w in model["weeks"]] == [(2, 1), (0, 3), (2, 0)]
    titles = [a["title"] for w in model["weeks"] for kind in ("assignments", "discussions") for a in w[kind]]
    assert len(titles) == len(set(titles)) == 8
    assert model["grade_plan"]["total_computed"] == 60
    assert sorted(r["name"] for r in model["rubrics"]) == ["Sample Assignment Rubric", "Sample Discussion Rubric"]
    assert not [d for d in resolve(model)["diagnostics"] if d["severity"] == "error"]


def test_zero_activities_is_a_valid_scaffold_without_unused_rubrics(tmp_path):
    model = read_model(generate(tmp_path, weeks=1, assignments_per_week=0, discussions_per_week=0,
                               worked_example=True))
    assert model["rubrics"] == []
    assert model["grade_plan"]["total_computed"] == 0
    assert model["weeks"][0]["assignments"] == []
    assert model["weeks"][0]["discussions"] == []


@pytest.mark.parametrize("kwargs", [
    {"weeks": 0}, {"weeks": 53}, {"weeks": True},
    {"assignments_per_week": -1}, {"discussions_per_week": 1.5},
    {"week_counts": {0: {"assignments": 2}}},
    {"week_counts": {1: {"quizzes": 1}}},
    {"week_counts": {1: {"discussions": True}}},
])
def test_invalid_scaffold_options_do_not_create_a_file(tmp_path, kwargs):
    with pytest.raises(ValueError):
        generate(tmp_path, **kwargs)
    assert not list(tmp_path.iterdir())


def test_writing_areas_use_normal_black_and_native_objective_bullets(tmp_path):
    doc = Document(generate(tmp_path, weeks=1))
    paragraphs = [p for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs]
    overview = next(p for p in paragraphs if p.text == WRITING_PROMPTS["week_introduction"])
    objective = next(p for p in paragraphs if p.text == WRITING_PROMPTS["learning_objective"])
    for para in (overview, objective):
        assert para.runs[0].italic is False
        assert para.runs[0].bold is False
        assert str(para.runs[0].font.color.rgb) == "000000"
    assert objective.style.name == "List Bullet"
    assert objective._p.xpath("./w:pPr/w:numPr")
    model = read_model(tmp_path / "blueprint.docx")
    assert next(b for b in model["weeks"][0]["objectives"] if b["kind"] == "placeholder")["list_item"] is True


def test_tables_repeat_headers_and_long_content_rows_can_split(tmp_path):
    doc = Document(generate(tmp_path, weeks=2, worked_example=True))
    grade = next(t for t in doc.tables if t.cell(0, 0).text == "Assignment")
    assert grade.rows[0]._tr.xpath("./w:trPr/w:tblHeader")
    assert grade.columns[0].width > grade.columns[1].width * 3
    for rubric in [t for t in doc.tables if t.cell(0, 0).text == "Criterion"]:
        assert rubric.rows[0]._tr.xpath("./w:trPr/w:tblHeader")
        assert not rubric.rows[1]._tr.xpath("./w:trPr/w:cantSplit")
    week = next(t for t in doc.tables if t.cell(0, 0).text == "Overview:")
    assert week.rows[0]._tr.xpath("./w:trPr/w:cantSplit")
    assert not week.rows[1]._tr.xpath("./w:trPr/w:cantSplit")
    assert week.cell(0, 0).paragraphs[0].paragraph_format.keep_with_next


def test_generated_guidance_is_accounted_for_and_never_becomes_checklist_content(tmp_path):
    model = read_model(generate(tmp_path, weeks=2))
    assert model["source_accounting"]["complete"]
    assert not [d for d in model["diagnostics"] if d["severity"] == "error"]
    assert all(w["checklist"]["items"] == [] for w in model["weeks"])
    assert any(e["disposition"] == "authoring_guidance" for e in model["source_accounting"]["entries"])


def test_copyable_examples_work_as_separate_paragraphs_in_a_week(tmp_path):
    from copy import deepcopy
    from docx.text.paragraph import Paragraph
    from blueprint_packager.build import Builder
    from blueprint_packager.verify import verify_build
    from test_pipeline import SETTINGS, REPO_ROOT, model_with_joins

    source = generate(tmp_path, weeks=1)
    doc = Document(source)
    week = next(t for t in doc.tables if t.cell(0, 0).text == 'Overview:')
    # An author replaces a complete natural prompt and continues with paragraphs.
    prompt = next(p for p in week.cell(1, 0).paragraphs if p.text == WRITING_PROMPTS['week_introduction'])
    prompt.text = 'This week compares two synthetic approaches.'
    week.cell(1, 0).add_paragraph('A second paragraph provides the connection.')

    paragraphs = doc.paragraphs
    first = next(i for i, p in enumerate(paragraphs) if p.text == 'Assignment: Week 2 Assignment: Source comparison')
    sample = paragraphs[first:first + 4]
    assert [p.text for p in sample][-1] == WRITING_PROMPTS['assignment_instructions']
    title = 'Week 1 Assignment: Source comparison'
    for paragraph in sample:
        element = deepcopy(paragraph._p)
        week.cell(7, 0)._tc.append(element)
        copied = Paragraph(element, week.cell(7, 0))
        copied.text = copied.text.replace('Week 2 Assignment: Source comparison', title)
        if copied.text == WRITING_PROMPTS['assignment_instructions']:
            copied.text = 'Compare two synthetic sources and submit a paragraph.'
    grade = next(t for t in doc.tables if t.cell(0, 0).text == 'Assignment')
    total = grade.rows[-1]
    row = grade.add_row()
    row.cells[0].text, row.cells[1].text = title, '20'
    total._tr.addprevious(row._tr)
    total.cells[1].text = '20'

    first = next(i for i, p in enumerate(paragraphs) if p.text == '[dropdown: Additional detail]')
    panel = paragraphs[first:first + 3]
    assert [p.text for p in panel] == ['[dropdown: Additional detail]', WRITING_PROMPTS['panel_content'], '[/dropdown]']
    for paragraph in panel:
        element = deepcopy(paragraph._p)
        week.cell(1, 0)._tc.append(element)
        copied = Paragraph(element, week.cell(1, 0))
        if copied.text == WRITING_PROMPTS['panel_content']:
            copied.text = 'Synthetic detail placed inside the expandable section.'
    doc.save(source)

    model = model_with_joins(source)
    assert not [d for d in model['diagnostics'] if d['severity'] == 'error']
    assert len(model['weeks'][0]['assignments']) == 2
    assert model['grade_plan']['total_computed'] == 20
    blocks = model['weeks'][0]['overview_page']['blocks']
    assert [b['kind'] for b in blocks] == ['p', 'p', 'dropdown']
    out = tmp_path / 'copy-workflow-build'
    Builder(model, {**SETTINGS, 'build_mode': 'preview'}, REPO_ROOT).run(out, force=False)
    assert verify_build(out, model) == 0
