"""Readable template prompts and reference examples have explicit dispositions."""
from __future__ import annotations

import json

import pytest
from docx import Document

import synth
from blueprint_packager import vocabulary as V
from blueprint_packager.authoring import AUTHOR_NOTE_PREFIX, WRITING_PROMPTS, marker, readiness
from blueprint_packager.build import BuildError, Builder
from blueprint_packager.html_render import RenderContext, Template, render_blocks
from blueprint_packager.reader import read_model
from blueprint_packager.render_text import plain
from blueprint_packager.verify import verify_build
from test_pipeline import SETTINGS, REPO_ROOT, model_with_joins


@pytest.mark.parametrize("prompt", WRITING_PROMPTS.values())
def test_exact_readable_writing_prompts_remain_unfinished(prompt):
    assert marker(prompt) == ("placeholder", prompt)
    assert marker("  " + prompt + "  ") == ("placeholder", prompt)
    assert marker(prompt + " Then explain your choice.") is None
    assert marker(prompt[:-1]) is None
    assert marker(prompt.lower()) is None


def test_author_note_requires_an_explicit_whole_paragraph_prefix():
    assert AUTHOR_NOTE_PREFIX == "Author note:"
    assert marker("Author note: Copy a complete activity block.") == ("authoring", "Copy a complete activity block.")
    assert marker("Author note: First line.\nSecond line.") == ("authoring", "First line.\nSecond line.")
    assert marker("The Author note: example is discussed here.") is None
    assert marker("Author notes are part of this lesson.") is None
    assert marker("Author note without a colon.") is None


@pytest.mark.parametrize("text,expected", [
    ("[authoring: Keep as editor guidance]", ("authoring", "Keep as editor guidance")),
    ("[placeholder: Write content]", ("placeholder", "Write content")),
    ("[video placeholder: Demonstration]", ("video placeholder", "Demonstration")),
    ("[video: Introduction | https://youtu.be/abcdefghijk]", ("video", "Introduction | https://youtu.be/abcdefghijk")),
])
def test_legacy_bracket_markers_still_work(text, expected):
    assert marker(text) == expected


def test_notes_excluded_prompts_pending_and_similar_prose_preserved(tmp_path):
    note = "Author note: This synthetic editing guidance stays out of the course."
    prompt = WRITING_PROMPTS["week_introduction"]
    prose = prompt + " Then connect the examples to your own experience."
    source = synth.canonical(tmp_path / "source.docx", overview=[
        "Page: Week 1 Overview", note, "Page heading: Week 1", prompt, prose])
    original = source.read_bytes()
    model = model_with_joins(source)
    assert source.read_bytes() == original
    blocks = model["weeks"][0]["overview_page"]["blocks"]
    assert not [d for d in model["diagnostics"] if d["severity"] == "error"]
    assert blocks[0]["kind"] == "placeholder" and blocks[0]["text"] == prompt
    assert plain(blocks[1]["runs"]) == prose
    entries = model["source_accounting"]["entries"]
    assert next(entry for entry in entries if entry["text"] == note)["disposition"] == "authoring_guidance"
    assert next(entry for entry in entries if entry["text"] == prompt)["disposition"] == "content"
    assert not readiness(model)["ready"]
    markup = render_blocks(blocks, RenderContext(Template.load("composite-navy"), "page"))
    assert prose in markup and note not in markup
    with pytest.raises(BuildError, match="draft"):
        Builder(model, SETTINGS, REPO_ROOT).run(tmp_path / "ready", force=False)


@pytest.mark.parametrize("frame_level", [1, 2])
def test_authoring_reference_samples_never_become_course_content(tmp_path, frame_level):
    source = synth.canonical(tmp_path / "reference.docx", frame_level=frame_level)
    doc = Document(source)
    doc.add_paragraph(V.H_AUTHORING_REFERENCE, style=f"Heading {frame_level}")
    doc.add_paragraph("Copy this complete example", style=f"Heading {frame_level + 1}")
    doc.add_paragraph("Assignment: Copyable assignment example")
    doc.add_paragraph("Points: 999 pts")
    doc.add_paragraph(WRITING_PROMPTS["assignment_instructions"])
    doc.add_paragraph("[video placeholder: Reference-only video]")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Week 99 - A sample that must not become a week"
    table.cell(0, 1).text = "Discussion: Copyable discussion example"
    table.cell(1, 0).text = "[dropdown: A sample]"
    table.cell(1, 1).text = "An intentionally incomplete reference sample."
    # A recognized next frame section ends the reference region.
    doc.add_paragraph(V.H_NOTES, style=f"Heading {frame_level}")
    doc.add_paragraph("A later extraction note remains a note.")
    doc.save(source)
    model = model_with_joins(source)
    assert not [d for d in model["diagnostics"] if d["severity"] == "error"]
    assert len(model["weeks"]) == 1
    assert len(model["weeks"][0]["assignments"]) == 1
    assert len(model["weeks"][0]["discussions"]) == 1
    assert readiness(model)["ready"]
    entries = model["source_accounting"]["entries"]
    reference_entries = [entry for entry in entries if entry["where"] == V.H_AUTHORING_REFERENCE]
    assert len(reference_entries) == 10
    assert all(entry["disposition"] == "authoring_guidance" for entry in reference_entries)
    assert next(entry for entry in entries if entry["text"] == "A later extraction note remains a note.")["disposition"] == "notes"
    if frame_level == 1:
        output = tmp_path / "build"
        Builder(model, SETTINGS, REPO_ROOT).run(output, force=False)
        assert verify_build(output, model) == 0
        packaged_text = "\n".join(path.read_text() for path in (output / "package").rglob("*")
                                  if path.is_file() and path.suffix in (".html", ".xml"))
        assert "Copyable assignment example" not in packaged_text
        assert "Reference-only video" not in packaged_text
        assert "Week 99" not in packaged_text
        assert json.loads((output / "readiness_report.json").read_text())["state"] == "ready"
