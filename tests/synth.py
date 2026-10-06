"""Synthetic blueprint documents for tests. No real course text appears here."""
from __future__ import annotations

from pathlib import Path

from docx import Document

from blueprint_packager import vocabulary as V

LABELS = {key: label for key, _prefix, label in V.SECTION_LABELS}


def _fill(cell, lines):
    first = True
    for line in lines:
        style = None
        if isinstance(line, tuple):
            line, style = line
        if first:
            para = cell.paragraphs[0]
            para.text = line
            first = False
        else:
            para = cell.add_paragraph(line)
        if style:
            para.style = style


def _one_column(doc, rows):
    table = doc.add_table(rows=len(rows), cols=1)
    for index, lines in enumerate(rows):
        _fill(table.rows[index].cells[0], lines)
    return table


def _front(doc, grade_rows):
    doc.add_paragraph("Blueprint", style="Title")
    doc.add_paragraph("SYN 101 Synthetic Course Blueprint")
    _one_column(doc, [["COURSE DESCRIPTION (catalog)"], ["A synthetic course used only for tests."]])
    _one_column(doc, [["TEXTBOOK/S OR REQUIRED MATERIALS"], ["No materials."]])
    _one_column(doc, [["COURSE LEARNING OUTCOMES"], ["Outcome one.", "Outcome two."]])
    table = doc.add_table(rows=len(grade_rows) + 2, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Assignment", "Points"
    for index, (name, points) in enumerate(grade_rows, start=1):
        table.rows[index].cells[0].text, table.rows[index].cells[1].text = name, str(points)
    table.rows[-1].cells[0].text = "Total Points"
    table.rows[-1].cells[1].text = str(sum(points for _name, points in grade_rows))


def _rubric(doc, name, used_by, heading=True):
    if heading:
        doc.add_paragraph(name, style="Heading 3")
    else:
        doc.add_paragraph(name)
    for user in used_by:
        doc.add_paragraph(f"Used by: {user}")
    table = doc.add_table(rows=3, cols=3)
    for col, text in enumerate(["Criteria", "Strong — 100%", "Weak — 0%"]):
        table.rows[0].cells[col].text = text
    for row, (criterion, strong, weak) in enumerate([("Clarity — 60%", "Clear throughout.", "Unclear."), ("Evidence — 40%", "Well supported.", "Unsupported.")], start=1):
        table.rows[row].cells[0].text, table.rows[row].cells[1].text, table.rows[row].cells[2].text = criterion, strong, weak
    return table


DEFAULT_OVERVIEW = ["Page: Week 1 Overview", "Page heading: Week 1 – Sample Topic", "Overview paragraph one.",
                    ("Key Ideas:", "Heading 3"),
                    "Dropdown / expandable section: First Idea", "Inside the first panel.",
                    "Checklist: this line is ordinary content, not a checklist object.",
                    "End of dropdown / expandable section",
                    "Visual cue: Callout", "Inside the callout.", "End of visual cue",
                    "Closing overview paragraph."]
DEFAULT_PRE_WEEK = [
    ["Getting Started"], ["Module: Getting Started"],
    ["Staff Notes"], ["Module: Staff Notes", "Visibility: Hidden from students"],
    ["Help Desk"], ["Link: Help Desk", "URL: https://example.org/help", "Placement: Staff Notes", "Visibility: Hidden from students"],
]


def canonical(path: Path, *, rubric_on_assignment: str = "Report Rubric", assignment_points: str = "60 pts",
              checklist: list[str] | None = None, overview: list | None = None, pre_week: list | None = None,
              assignment_body: list | None = None) -> Path:
    """A blueprint already in the canonical layout."""
    doc = Document()
    _front(doc, [("Week 1 Assignment: Report", 60), ("Week 1 Discussion: Opening", 40)])
    doc.add_paragraph(V.H_COURSE_INTRO, style="Heading 2")
    doc.add_paragraph("Page: Welcome to SYN 101")
    doc.add_paragraph("Page heading: Welcome to SYN 101")
    doc.add_paragraph("Placement: Getting Started")
    doc.add_paragraph("This is the welcome text.")
    doc.add_paragraph(V.H_BEFORE_WEEK_1, style="Heading 2")
    _one_column(doc, pre_week or DEFAULT_PRE_WEEK)
    doc.add_paragraph("Week 1 - Sample Topic", style="Heading 2")
    _one_column(doc, [
        [LABELS["overview"]],
        overview or DEFAULT_OVERVIEW,
        [LABELS["objectives"], V.OBJECTIVES_SUBLABEL],
        ["By the end of this week, students will be able to:", "Do the first thing.", "Do the second thing."],
        [LABELS["resources"]],
        [("Required Resources:", "Heading 3"), "Author, A. (2026). A synthetic reading."],
        [LABELS["assignments"]],
        ["Assignment: Week 1 Assignment: Report", f"Points: {assignment_points}", "Gradebook item: Week 1 Assignment: Report",
         f"Rubric: {rubric_on_assignment}", *(assignment_body or [("The Report", "Heading 3"), "Write a short synthetic report.",
                                                                  ("Submission Requirements", "Heading 4"), "One document."])],
        [LABELS["discussions"]],
        ["Discussion: Week 1 Discussion: Opening", "Points: 40 pts", "Gradebook item: Week 1 Discussion: Opening",
         "Rubric: Discussion Rubric", "Forum: Discussions", "Introduce yourself to the synthetic class."],
        [LABELS["checklist"]],
        checklist or ["Checklist: Week 1 Checklist", "Post to the discussion", "Submit the report"],
    ])
    doc.add_paragraph(V.H_COURSE_ACTIVITIES, style="Heading 2")
    doc.add_paragraph("Discussion: Ask a Question")
    doc.add_paragraph("Forum: Community Forum")
    doc.add_paragraph("Post course questions here.")
    doc.add_paragraph(V.H_RUBRICS, style="Heading 2")
    _rubric(doc, "Report Rubric", ["Assignment: Week 1 Assignment: Report"])
    _rubric(doc, "Discussion Rubric", ["Discussion: Week 1 Discussion: Opening"])
    doc.add_paragraph(V.H_NOTES, style="Heading 2")
    doc.add_paragraph("Synthetic note.")
    doc.save(path)
    return path


def messy(path: Path) -> Path:
    """A wizard-layout blueprint after hand editing: misplaced objects and stale wizard lines."""
    doc = Document()
    _front(doc, [("Week 1 Assignment: Report", 60), ("Week 1 Discussion: Opening", 40)])
    doc.add_paragraph(V.H_COURSE_INTRO, style="Heading 2")
    doc.add_paragraph("Welcome to SYN 101")
    doc.add_paragraph("This is the welcome text.")
    doc.add_paragraph(V.H_BEFORE_WEEK_1, style="Heading 2")
    _one_column(doc, [["Staff Notes"], ["Hidden manifest item: Object: Staff Notes; type: module; hidden."]])
    doc.add_paragraph("Course Content:")
    doc.add_paragraph("Week 1 - Sample Topic", style="Heading 2")
    _one_column(doc, [
        [LABELS["overview"]],
        ["Week 1 Overview:", "Old HTML Title", "Week 1 – Sample Topic:", "Overview paragraph one.",
         "Dropdown / expandable section: First Idea", "Inside the first panel.", "After the panel.",
         "Reading typed into the overview by mistake."],
        [LABELS["objectives"], V.OBJECTIVES_SUBLABEL],
        ["By the end of this week, students will be able to:", "Do the first thing."],
        [LABELS["resources"]],
        [""],
        [LABELS["assignments"]],
        [("Opening Discussion", "Heading 2"), "Introduce yourself to the synthetic class.",
         "Old Report Exercise:", "Points: 1 pts", "Rubric: 9: Old Rubric",
         ("The Report", "Heading 2"), "Write a short synthetic report.",
         ("A body sentence that was styled as a heading by accident.", "Heading 3")],
        [LABELS["discussions"]],
        [""],
        [LABELS["checklist"]],
        ["Week 1 Checklist:", "Post to the discussion"],
    ])
    doc.add_paragraph(V.H_COURSE_ACTIVITIES, style="Heading 2")
    doc.add_paragraph("Ask a Question:")
    doc.add_paragraph("Post course questions here.")
    doc.add_paragraph(V.H_RUBRICS, style="Heading 2")
    _rubric(doc, "Report Rubric", [], heading=False)
    _rubric(doc, "Discussion Rubric", [], heading=False)
    doc.add_paragraph(V.H_NOTES, style="Heading 2")
    doc.add_paragraph("Synthetic note.")
    doc.save(path)
    return path
