import json
import zipfile
from pathlib import Path

import pytest

import synth
from blueprint_packager import identity as I
from blueprint_packager import vocabulary as V
from blueprint_packager.docx_ingest import Para, Table, read_docx
from blueprint_packager.fidelity import prove
from blueprint_packager.joins import resolve
from blueprint_packager.normalize import MapError, Normalizer, normalize
from blueprint_packager.reader import read_model

REPO_ROOT = Path(__file__).resolve().parents[3]
HAS_WORKBENCH = (REPO_ROOT / "scripts" / "d2l_package_lib.py").is_file()


def model_with_joins(path):
    model = read_model(path)
    resolved = resolve(model)
    model["joins"] = resolved["joins"]
    model["diagnostics"] = model["diagnostics"] + resolved["diagnostics"]
    return model


def errors(model):
    return [d for d in model["diagnostics"] if d["severity"] == "error"]


# ----------------------------------------------------------------- vocabulary
def test_directive_parsing_is_keyword_and_colon_only():
    assert V.parse_directive("Assignment: Week 1 Assignment: Report") == ("Assignment", "Week 1 Assignment: Report")
    assert V.parse_directive("Points: 10 pts") == ("Points", "10 pts")
    assert V.parse_directive("End of dropdown / expandable section") == (V.D_DROPDOWN_END, "")
    assert V.parse_directive("Rubric ID: 2 | Resource code: X") is None
    assert V.parse_directive("Points are awarded for clarity.") is None
    assert V.parse_points("10 pts") == 10 and V.parse_points("7.5") == 7.5 and V.parse_points("many") is None


def test_identity_keys_match_the_workbench_construction():
    if not HAS_WORKBENCH:
        pytest.skip("Workbench scripts not present")
    import sys

    sys.path.insert(0, str(REPO_ROOT / "scripts"))
    import course_artifact_contracts as contracts

    assert I.scoped_key("entity", "a", "b", "c") == contracts._scoped_key("entity", "a", "b", "c")
    assert I.canonical_sha256({"b": 1, "a": [2]}) == contracts.canonical_sha256({"b": 1, "a": [2]})


# --------------------------------------------------------------------- reader
def test_reader_builds_the_course_model(tmp_path):
    model = model_with_joins(synth.canonical(tmp_path / "canonical.docx"))
    assert errors(model) == []
    assert [m["title"] for m in model["modules"]] == ["Getting Started", "Staff Notes"]
    assert model["links"][0]["url"] == "https://example.org/help"
    week = model["weeks"][0]
    assert week["overview_page"]["title"] == "Week 1 Overview"
    assert week["overview_page"]["page_heading"] == "Week 1 – Sample Topic"
    assert [a["title"] for a in week["assignments"]] == ["Week 1 Assignment: Report"]
    assert week["assignments"][0]["points"] == 60 and week["assignments"][0]["rubric"] == "Report Rubric"
    assert week["discussions"][0]["forum"] == "Discussions"
    assert week["checklist"]["title"] == "Week 1 Checklist" and len(week["checklist"]["items"]) == 2
    assert model["grade_plan"]["total_computed"] == 100
    rubric = model["rubrics"][0]
    assert [level["percent"] for level in rubric["grid"]["levels"]] == [100.0, 0.0]
    assert [c["weight"] for c in rubric["grid"]["criteria"]] == [60.0, 40.0]


def test_keywords_open_objects_only_where_that_object_can_live(tmp_path):
    model = model_with_joins(synth.canonical(tmp_path / "canonical.docx"))
    blocks = model["weeks"][0]["overview_page"]["blocks"]
    dropdown = next(b for b in blocks if b["kind"] == "dropdown")
    texts = ["".join(r["text"] for r in b["runs"]) for b in dropdown["blocks"]]
    assert "Checklist: this line is ordinary content, not a checklist object." in texts
    assert any(b["kind"] == "visual" and b["cue"] == "Callout" for b in blocks)


def test_broken_rubric_join_is_an_error_not_a_guess(tmp_path):
    model = model_with_joins(synth.canonical(tmp_path / "broken.docx", rubric_on_assignment="Report Rubrik"))
    codes = {d["code"] for d in errors(model)}
    assert "RUBRIC_UNKNOWN" in codes
    assert "RUBRIC_USER_UNKNOWN" not in codes and "RUBRIC_NOT_RECIPROCAL" in codes


# ----------------------------------------------------------------- normalizer
def course_map_for(source: Path) -> dict:
    doc = read_docx(source)

    def pid(text: str) -> str:
        return next(p.id for p in doc.paras() if p.text().strip() == text)

    tables = list(doc.tables())
    front = [t.id for t in tables[:4]]
    pre, week = tables[4], tables[5]
    rubric_tables = [t.id for t in tables[6:8]]
    cells = [row[0] for row in week.rows]
    return {
        "schema": "blueprint_packager.course_map/1",
        "source": {"sha256": doc.sha256},
        "policy": {"suggestions": "proposed"},
        "title": [pid("Blueprint"), pid("SYN 101 Synthetic Course Blueprint")],
        "front_tables": front, "grade_plan_table": front[3],
        "label_templates": {
            "overview": [pid(synth.LABELS["overview"])], "objectives": [pid(synth.LABELS["objectives"]), pid(V.OBJECTIVES_SUBLABEL)],
            "resources": [pid(synth.LABELS["resources"])], "assignments": [pid(synth.LABELS["assignments"])],
            "discussions": [pid(synth.LABELS["discussions"])], "checklist": [pid(synth.LABELS["checklist"])],
        },
        "course_introduction": {"heading": pid(V.H_COURSE_INTRO), "objects": [{
            "type": "page", "title": {"carry": pid("Welcome to SYN 101")},
            "directives": [["Page heading", {"carry": pid("Welcome to SYN 101")}], ["Placement", {"text": "Staff Notes", "source": "test"}]],
            "content": [pid("This is the welcome text.")]}]},
        "before_week_1": {"heading": pid(V.H_BEFORE_WEEK_1), "template_table": pre.id, "rows": [{
            "label": {"carry": pid("Staff Notes")},
            "objects": [{"type": "module", "title": {"carry": pid("Staff Notes")},
                         "supersede": [pid("Hidden manifest item: Object: Staff Notes; type: module; hidden.")],
                         "directives": [["Visibility", {"text": "Hidden from students", "source": "wizard line"}]]}]}]},
        "course_content_line": pid("Course Content:"),
        "weeks": [{
            "number": 1, "template_table": week.id, "heading": {"keep": pid("Week 1 - Sample Topic")},
            "labels": [pid(synth.LABELS[k]) for k in V.SECTION_KEYS] + [pid(V.OBJECTIVES_SUBLABEL)],
            "overview": {"page": {"title": {"carry": pid("Week 1 Overview:"), "strip_colon": True},
                                  "directives": [["Page heading", {"carry": pid("Week 1 – Sample Topic:"), "strip_colon": True}]]},
                         "content": [f"{pid('Overview paragraph one.')}:{pid('After the panel.')}"]},
            "objectives": {"content": [f"@{cells[3].id}"]},
            "resources": {"content": [pid("Reading typed into the overview by mistake.")]},
            "assignments": [{"type": "assignment", "grade_item": "Week 1 Assignment: Report", "rubric": "Report Rubric",
                             "supersede": [f"{pid('Old Report Exercise:')}:{pid('Rubric: 9: Old Rubric')}"],
                             "content": [f"{pid('The Report')}:{pid('A body sentence that was styled as a heading by accident.')}"]}],
            "discussions": [{"type": "discussion", "grade_item": "Week 1 Discussion: Opening", "rubric": "Discussion Rubric",
                             "directives": [["Forum", {"text": "Discussions", "source": "test"}]],
                             "content": [f"{pid('Opening Discussion')}:{pid('Introduce yourself to the synthetic class.')}"]}],
            "checklist": {"title": {"carry": pid("Week 1 Checklist:"), "strip_colon": True}, "content": [pid("Post to the discussion")]},
        }],
        "course_activities": {"heading": pid(V.H_COURSE_ACTIVITIES), "objects": [{
            "type": "discussion", "title": {"carry": pid("Ask a Question:"), "strip_colon": True},
            "directives": [["Forum", {"text": "Community Forum", "source": "test"}]],
            "content": [pid("Post course questions here.")]}]},
        "rubrics": {"heading": pid(V.H_RUBRICS), "blocks": [
            {"name": pid("Report Rubric"), "table": rubric_tables[0], "used_by": ["Assignment: Week 1 Assignment: Report"]},
            {"name": pid("Discussion Rubric"), "table": rubric_tables[1], "used_by": ["Discussion: Week 1 Discussion: Opening"]}]},
        "notes": {"heading": pid(V.H_NOTES), "content": [pid("Synthetic note.")]},
        "dropdowns": {pid("Dropdown / expandable section: First Idea"): pid("Inside the first panel.")},
        "supersede": [{"ids": [pid("Old HTML Title")], "reason": "wizard html title"}],
        "restyle": {pid("Opening Discussion"): "heading 3", pid("The Report"): "heading 3",
                    pid("A body sentence that was styled as a heading by accident."): "Normal"},
    }


def test_normalization_moves_structure_and_keeps_every_word(tmp_path):
    source = synth.messy(tmp_path / "messy.docx")
    map_path = tmp_path / "map.json"
    map_path.write_text(json.dumps(course_map_for(source)), encoding="utf-8")
    out = tmp_path / "normalized.docx"
    result = normalize(source, map_path, out)
    report = prove(result, out)
    assert report["passed"], [c for c in report["checks"] if not c["passed"]]
    model = model_with_joins(out)
    assert errors(model) == []
    week = model["weeks"][0]
    assert [d["title"] for d in week["discussions"]] == ["Week 1 Discussion: Opening"]  # moved out of the assignment cell
    assert week["assignments"][0]["points"] == 60  # stale "Points: 1 pts" superseded by the points table
    assert any("Reading typed into the overview by mistake." in "".join(r["text"] for r in b.get("runs", [])) for b in week["resources"])
    dispositions = {d.disposition for d in result["dispositions"].values()}
    assert {"structure_superseded", "carried_into_directive", "kept_restyled"} <= dispositions


def test_an_unmapped_paragraph_stops_the_run(tmp_path):
    source = synth.messy(tmp_path / "messy.docx")
    cmap = course_map_for(source)
    cmap["notes"]["content"] = []  # leaves "Synthetic note." unaccounted for
    map_path = tmp_path / "map.json"
    map_path.write_text(json.dumps(cmap), encoding="utf-8")
    with pytest.raises(MapError, match="not accounted for"):
        normalize(source, map_path, tmp_path / "out.docx")


def test_a_map_written_for_another_file_is_refused(tmp_path):
    source = synth.messy(tmp_path / "messy.docx")
    cmap = course_map_for(source)
    cmap["source"]["sha256"] = "0" * 64
    map_path = tmp_path / "map.json"
    map_path.write_text(json.dumps(cmap), encoding="utf-8")
    with pytest.raises(MapError, match="written for source"):
        normalize(source, map_path, tmp_path / "out.docx")


def _normalized(tmp_path, cmap, source, name="out.docx"):
    map_path = tmp_path / f"{name}.map.json"
    map_path.write_text(json.dumps(cmap), encoding="utf-8")
    out = tmp_path / name
    result = normalize(source, map_path, out)
    report = prove(result, out)
    assert report["passed"], [c for c in report["checks"] if not c["passed"]]
    return result, report, model_with_joins(out)


def _pid(source, text):
    return next(p.id for p in read_docx(source).paras() if p.text().strip() == text)


def test_a_title_can_follow_the_activitys_own_heading(tmp_path):
    source = synth.messy(tmp_path / "messy.docx")
    cmap = course_map_for(source)
    heading = _pid(source, "The Report")
    cmap["weeks"][0]["assignments"][0]["title"] = {"copy": heading}
    cmap["rubrics"]["blocks"][0]["used_by"] = [{"activity": "Assignment", "copy": heading}]
    result, _, model = _normalized(tmp_path, cmap, source)
    assignment = model["weeks"][0]["assignments"][0]
    assert assignment["title"] == "The Report"
    assert assignment["grade_item"] == "Week 1 Assignment: Report"  # the points table row stays the join
    assert result["dispositions"][heading].disposition != "carried_into_directive"  # the heading stays in the instructions
    assert errors(model) == []
    assert "TITLE_DIFFERS_FROM_GRADE_ROW" in {d["code"] for d in model["diagnostics"]}


def test_points_come_from_the_body_and_fall_back_to_the_points_table(tmp_path):
    source = synth.messy(tmp_path / "messy.docx")
    line = _pid(source, "Points: 1 pts")

    undecided = course_map_for(source)
    undecided["weeks"][0]["assignments"][0]["body_points"] = {"line": line}
    path = tmp_path / "undecided.json"
    path.write_text(json.dumps(undecided), encoding="utf-8")
    with pytest.raises(MapError, match="must say which"):
        normalize(source, path, tmp_path / "undecided.docx")

    table_wins = course_map_for(source)
    table_wins["weeks"][0]["assignments"][0]["body_points"] = {"line": line, "use": "points_table", "reason": "stale"}
    result, report, model = _normalized(tmp_path, table_wins, source, "table.docx")
    assert model["weeks"][0]["assignments"][0]["points"] == 60
    assert report["points_decisions"][0]["conflict"] is True

    body_wins = course_map_for(source)
    body_wins["weeks"][0]["assignments"][0]["body_points"] = {"line": line, "use": "body"}
    _, _, model = _normalized(tmp_path, body_wins, source, "body.docx")
    assert model["weeks"][0]["assignments"][0]["points"] == 1
    assert errors(model) == []
    assert "POINTS_DIFFER_FROM_GRADE_ROW" in {d["code"] for d in model["diagnostics"] if d["severity"] == "warning"}


def test_authorized_additions_are_the_only_new_words(tmp_path):
    source = synth.messy(tmp_path / "messy.docx")
    anchor = _pid(source, "Introduce yourself to the synthetic class.")
    cmap = course_map_for(source)
    cmap["additions"] = [{
        "after": anchor, "authorized_by": "test operator", "source": "test comment thread",
        "paragraphs": [{"like": anchor, "runs": [{"text": "You may also use "}, {"text": "this example", "href": "https://example.org/a"}]}]}]
    result, report, model = _normalized(tmp_path, cmap, source)
    assert [a["text"] for a in result["additions"]] == ["You may also use this example"]
    assert report["authorized_additions"][0]["links"] == ["https://example.org/a"]
    blocks = model["weeks"][0]["discussions"][0]["blocks"]
    added = next(b for b in blocks if "You may also use" in "".join(r["text"] for r in b.get("runs", [])))
    assert [r["href"] for r in added["runs"] if r.get("href")] == ["https://example.org/a"]

    unsigned = course_map_for(source)
    unsigned["additions"] = [{"after": anchor, "paragraphs": [{"runs": [{"text": "Unattributed."}]}]}]
    path = tmp_path / "unsigned.json"
    path.write_text(json.dumps(unsigned), encoding="utf-8")
    with pytest.raises(MapError, match="authorized_by"):
        normalize(source, path, tmp_path / "unsigned.docx")


def _edited_copy(source: Path, target: Path, before_everything=False) -> Path:
    """The same blueprint after its author added two paragraphs and removed one."""
    import copy as _copy

    from docx import Document as _Document
    from docx.oxml.ns import qn

    doc = _Document(source)
    paragraphs = {"".join(t.text or "" for t in p.iter(qn("w:t"))).strip(): p for p in doc.element.body.iter(qn("w:p"))}

    def add_after(anchor_text: str, text: str) -> None:
        anchor = paragraphs[anchor_text]
        new = _copy.deepcopy(anchor)
        texts = list(new.iter(qn("w:t")))
        texts[0].text = text
        for extra in texts[1:]:
            extra.text = ""
        anchor.addnext(new)

    def add_before(anchor_text: str, text: str) -> None:
        anchor = paragraphs[anchor_text]
        new = _copy.deepcopy(anchor)
        texts = list(new.iter(qn("w:t")))
        texts[0].text = text
        for extra in texts[1:]:
            extra.text = ""
        anchor.addprevious(new)

    add_after("This is the welcome text.", "A new welcome sentence.")  # between two mapped stretches
    add_after("Overview paragraph one.", "A new overview paragraph.")  # inside a mapped range
    if before_everything:
        add_before("Blueprint", "A line typed above the title.")  # before anything the map names
    stale = paragraphs["Old HTML Title"]
    stale.getparent().remove(stale)
    doc.save(target)
    return target


def test_a_map_follows_the_document_when_paragraphs_shift(tmp_path):
    from blueprint_packager import rebase

    source = synth.messy(tmp_path / "messy.docx")
    later = _edited_copy(source, tmp_path / "later.docx")
    old, new = read_docx(source), read_docx(later)
    alignment = rebase.align(old, new)
    assert alignment["counts"]["added"] == 2 and alignment["counts"]["removed"] == 1
    assert not alignment["identity"]
    rebased, problems = rebase.rebase_map(course_map_for(source), alignment, new, old)
    assert rebased["source"]["sha256"] == new.sha256
    assert [p for p in problems if p.get("severity") != "note"] == []  # the removed line was only in a list of lines to drop

    # a paragraph added between two mapped stretches joins the content stretch before it, and the
    # placement is reported as an inference; the run does not stop (owner's rule, 2026-10-07)
    welcome, new_line = _pid(later, "This is the welcome text."), _pid(later, "A new welcome sentence.")
    assert rebased["course_introduction"]["objects"][0]["content"] == [f"{welcome}:{new_line}"]
    placed = alignment["inferred_placements"]
    assert [p["new"] for p in placed] == [new_line] and placed[0]["joined"] == f"{welcome}:{new_line}"
    assert "course_introduction" in placed[0]["where"]
    rep = rebase.report(alignment, problems, old, new)
    assert rep["placed_by_inference"] == placed
    assert "## Placed by inference, confirm" in rebase.report_markdown(rep)

    # the run passes as rebased; the paragraph added inside a range came along by itself, as before
    _, _, model = _normalized(tmp_path, rebased, later, "later_normalized.docx")
    assert errors(model) == []
    overview = " ".join("".join(r["text"] for r in b.get("runs", [])) for b in model["weeks"][0]["overview_page"]["blocks"])
    assert "A new overview paragraph." in overview
    welcome_page = " ".join("".join(r["text"] for r in b.get("runs", [])) for b in model["pages"][0]["blocks"])
    assert "A new welcome sentence." in welcome_page


def test_a_new_paragraph_with_no_mapped_neighbour_still_stops_the_run(tmp_path):
    from blueprint_packager import rebase

    source = synth.messy(tmp_path / "messy.docx")
    later = _edited_copy(source, tmp_path / "later.docx", before_everything=True)
    old, new = read_docx(source), read_docx(later)
    alignment = rebase.align(old, new)
    assert alignment["counts"]["added"] == 3
    rebased, problems = rebase.rebase_map(course_map_for(source), alignment, new, old)
    stray = _pid(later, "A line typed above the title.")
    blocking = [p for p in problems if p.get("severity") != "note"]
    assert [p["id"] for p in blocking] == [stray] and "place it by hand" in blocking[0]["problem"]
    assert [p["new"] for p in alignment["inferred_placements"]] == [_pid(later, "A new welcome sentence.")]  # the other one still joins
    path = tmp_path / "rebased.json"
    path.write_text(json.dumps(rebased), encoding="utf-8")
    with pytest.raises(MapError, match="A line typed above the title."):
        normalize(later, path, tmp_path / "refused.docx")


def test_authorized_edits_are_applied_recorded_and_the_only_wording_changes(tmp_path):
    source = synth.messy(tmp_path / "messy.docx")
    cmap = course_map_for(source)
    cmap["edits"] = [{"authorized_by": "test operator", "reason": "test",
                      "drop": [_pid(source, "Introduce yourself to the synthetic class.")],
                      "replace": [{"para": _pid(source, "Write a short synthetic report."), "find": "short", "with": "brief"}]}]
    result, report, model = _normalized(tmp_path, cmap, source)
    kinds = sorted(edit["kind"] for edit in report["authorized_edits"])
    assert kinds == ["drop", "replace"]
    week = model["weeks"][0]
    body = " ".join("".join(r["text"] for r in b.get("runs", [])) for b in week["assignments"][0]["blocks"])
    assert "Write a brief synthetic report." in body and "short" not in body
    assert week["discussions"][0]["blocks"] == [] or all("Introduce yourself" not in "".join(r["text"] for r in b.get("runs", []))
                                                           for b in week["discussions"][0]["blocks"])
    dispositions = {d.disposition for d in result["dispositions"].values()}
    assert "dropped_authorized" in dispositions and any("edited" in d for d in dispositions)

    for bad, message in (
        ({"reason": "no name", "drop": [_pid(source, "Synthetic note.")]}, "authorized_by"),
        ({"authorized_by": "x", "reason": "twice", "replace": [{"para": _pid(source, "Write a short synthetic report."), "find": "t", "with": "T"}]}, "exactly once"),
    ):
        broken = course_map_for(source)
        broken["edits"] = [bad]
        path = tmp_path / "bad.json"
        path.write_text(json.dumps(broken), encoding="utf-8")
        with pytest.raises(MapError, match=message):
            normalize(source, path, tmp_path / "bad.docx")


def test_a_rubric_can_be_cloned_with_named_edits(tmp_path):
    source = synth.messy(tmp_path / "messy.docx")
    cmap = course_map_for(source)
    original = cmap["rubrics"]["blocks"][0]
    cmap["rubrics"]["blocks"].insert(1, {
        "clone": {"name": "Report Rubric (Short)", "table": original["table"],
                  "cell_edits": [{"find": "Clear throughout.", "with": "Clear."}],
                  "authorized_by": "test operator", "source": "copy of Report Rubric"},
        "used_by": ["Assignment: Week 1 Assignment: Report"]})
    original["used_by"] = []
    cmap["weeks"][0]["assignments"][0]["rubric"] = "Report Rubric (Short)"
    _, report, model = _normalized(tmp_path, cmap, source)
    assert errors(model) == []
    rubrics = {r["name"]: r for r in model["rubrics"]}

    def cell(name: str) -> str:
        return "".join(run["text"] for block in rubrics[name]["grid"]["criteria"][0]["cells"][0] for run in block["runs"])

    assert cell("Report Rubric (Short)") == "Clear." and cell("Report Rubric") == "Clear throughout."
    assert [e["kind"] for e in report["authorized_edits"]] == ["rubric_clone_edit"]
    assert model["weeks"][0]["assignments"][0]["rubric"] == "Report Rubric (Short)"


def test_a_title_inherits_its_week_only_when_it_lacks_one(tmp_path):
    source = synth.messy(tmp_path / "messy.docx")
    heading = _pid(source, "The Report")
    for prefix, expected in (("Week 1", "Week 1 The Report"), ("The", "The Report")):
        cmap = course_map_for(source)
        cmap["weeks"][0]["assignments"][0]["title"] = {"copy": heading, "prefix_if_missing": prefix}
        cmap["rubrics"]["blocks"][0]["used_by"] = [{"activity": "Assignment", "copy": heading, "prefix_if_missing": prefix}]
        _, _, model = _normalized(tmp_path, cmap, source, f"{prefix}.docx")
        assert model["weeks"][0]["assignments"][0]["title"] == expected
        assert errors(model) == []


def test_the_calendar_turns_a_week_and_a_day_into_a_timestamp():
    from blueprint_packager.calendar import Calendar, CalendarError
    from blueprint_packager.html_render import drop_query_parameters

    calendar = Calendar({"timezone": "America/New_York", "weeks": {"1": ["2026-10-28", "2026-11-01"], "8": ["2026-12-14", "2026-12-18"]},
                         "rules": {"assignment": {"day": "last_day"}}, "overrides": {"Special": {"assignment": {"day": "Friday"}}}})
    assert calendar.stamp(1, "Friday")["utc"] == "2026-10-31T03:59:00.000Z"   # daylight time
    assert calendar.stamp(1, "last_day")["utc"] == "2026-11-02T04:59:00.000Z"  # standard time, same week
    assert calendar.due("assignment", 8)["weekday"] == "Friday"                # the last week ends on a Friday
    assert calendar.due("assignment", 1, "Special")["weekday"] == "Friday"
    with pytest.raises(CalendarError, match="no Sunday"):
        calendar.day(8, "Sunday")
    calendar.rules["discussion_initial"] = {"day": "Friday", "last_week_day": "Thursday"}
    assert calendar.due("discussion_initial", 1)["weekday"] == "Friday"
    assert calendar.due("discussion_initial", 8)["weekday"] == "Thursday"  # the last week has its own day

    rules = [{"name": "utm_source", "value": "chatgpt.com"}]
    assert drop_query_parameters("https://example.org/a?utm_source=chatgpt.com", rules) == "https://example.org/a"
    assert drop_query_parameters("https://example.org/a?x=1&utm_source=chatgpt.com#top", rules) == "https://example.org/a?x=1#top"
    assert drop_query_parameters("https://example.org/a?utm_source=newsletter", rules) == "https://example.org/a?utm_source=newsletter"


# ---------------------------------------------------------------------- build
SETTINGS = {
    "schema": "blueprint_packager.build_settings/0.1",
    "package": {"title": "SYN 101 test build", "manifest_identifier": "D2L_SYN101", "base_resource_code": "00000000-0000-4000-8000-000000000101"},
    "template": "bumk515-classic",
    "forum_order": ["Community Forum", "Discussions"],
}


@pytest.mark.skipif(not HAS_WORKBENCH, reason="Workbench building blocks not present")
def test_build_is_byte_stable_and_reads_back(tmp_path):
    from blueprint_packager.build import Builder
    from blueprint_packager.verify import verify_build

    model = model_with_joins(synth.canonical(tmp_path / "canonical.docx"))
    digests = []
    for name in ("a", "b"):
        result = Builder(model, SETTINGS, REPO_ROOT).run(tmp_path / name, force=True)
        digests.append(I.sha256_file(result["zip"]))
    assert digests[0] == digests[1]
    with zipfile.ZipFile(tmp_path / "a" / "package.zip") as archive:
        names = set(archive.namelist())
    assert {"imsmanifest.xml", "grades_d2l.xml", "rubrics_d2l.xml", "dropbox_d2l.xml", "discussion_d2l_1.xml",
            "discussion_d2l_2.xml", "checklist_d2l.xml", "course_pages/week-1-overview.html"} <= names
    page = (tmp_path / "a" / "package" / "course_pages" / "week-1-overview.html").read_text(encoding="utf-8")
    assert "<details>" in page and "<summary>First Idea</summary>" in page and 'class="callout"' in page
    assert page.count("<h1>") == 1
    assert verify_build(tmp_path / "a", model) == 0


@pytest.mark.skipif(not HAS_WORKBENCH, reason="Workbench building blocks not present")
def test_build_refuses_a_model_with_errors(tmp_path):
    from blueprint_packager.build import BuildError, Builder

    model = model_with_joins(synth.canonical(tmp_path / "broken.docx", rubric_on_assignment="Report Rubrik"))
    with pytest.raises(BuildError, match="error"):
        Builder(model, SETTINGS, REPO_ROOT).run(tmp_path / "out", force=True)


@pytest.mark.skipif(not HAS_WORKBENCH, reason="Workbench building blocks not present")
@pytest.mark.parametrize("naming, expected", [("activity_title", "The Report"), ("points_table", "Week 1 Assignment: Report")])
def test_grade_items_are_named_by_the_chosen_source(tmp_path, naming, expected):
    from blueprint_packager.build import Builder
    from blueprint_packager.verify import verify_build

    source = synth.messy(tmp_path / "messy.docx")
    cmap = course_map_for(source)
    heading = _pid(source, "The Report")
    cmap["weeks"][0]["assignments"][0]["title"] = {"copy": heading}
    cmap["rubrics"]["blocks"][0]["used_by"] = [{"activity": "Assignment", "copy": heading}]
    _, _, model = _normalized(tmp_path, cmap, source)
    build_dir = tmp_path / naming
    Builder(model, {**SETTINGS, "grade_item_names": naming}, REPO_ROOT).run(build_dir, force=True)
    grades = (build_dir / "package" / "grades_d2l.xml").read_text(encoding="utf-8")
    assert expected in grades
    assert verify_build(build_dir, model) == 0


CALENDAR = {"timezone": "America/New_York", "weeks": {"1": ["2026-10-28", "2026-11-01"]},
            "rules": {"assignment": {"day": "last_day"}, "discussion_initial": {"day": "Friday"}, "discussion_response": {"day": "last_day"}}}


@pytest.mark.skipif(not HAS_WORKBENCH, reason="Workbench building blocks not present")
def test_a_checklist_can_be_built_from_its_weeks_activities_with_dates(tmp_path):
    import xml.etree.ElementTree as ET

    from blueprint_packager.build import Builder
    from blueprint_packager.verify import verify_build

    source = synth.canonical(tmp_path / "derived.docx", checklist=[
        "Checklist: Week 1 Checklist", f"{V.D_ITEMS}: {V.ITEMS_FROM_ACTIVITIES}", "Bring a notebook"])
    model = model_with_joins(source)
    assert errors(model) == [] and model["weeks"][0]["checklist"]["derive"] is True
    settings = {**SETTINGS, "calendar": CALENDAR, "checklists": {"extra_item_dates": {"Bring a notebook": {"day": "Wednesday"}}}}
    Builder(model, settings, REPO_ROOT).run(tmp_path / "out", force=True)
    items = [(item.findtext("name"), item.findtext("date_end")) for item in ET.parse(tmp_path / "out" / "package" / "checklist_d2l.xml").getroot().iter("item")]
    assert items == [
        ("Week 1 Discussion: Initial Post", "2026-10-31T03:59:00.000Z"),
        ("Week 1 Discussion: Response Posts", "2026-11-02T04:59:00.000Z"),
        ("Submit your Week 1 Assignment", None),          # the folder carries its own due date
        ("Bring a notebook", "2026-10-29T03:59:00.000Z"),  # an item listed in the blueprint, dated by a setting
    ]
    dropbox = (tmp_path / "out" / "package" / "dropbox_d2l.xml").read_text(encoding="utf-8")
    assert "<date_due>2026-11-02T04:59:00.000Z</date_due>" in dropbox
    assert verify_build(tmp_path / "out", model) == 0


@pytest.mark.skipif(not HAS_WORKBENCH, reason="Workbench building blocks not present")
def test_points_that_disagree_need_a_recorded_determination(tmp_path):
    from blueprint_packager.build import BuildError, Builder
    from blueprint_packager.verify import verify_build

    model = model_with_joins(synth.canonical(tmp_path / "points.docx", assignment_points="50 pts"))
    assert errors(model) == []
    assert "POINTS_DIFFER_FROM_GRADE_ROW" in {d["code"] for d in model["diagnostics"]}
    with pytest.raises(BuildError, match="no determination is recorded"):
        Builder(model, SETTINGS, REPO_ROOT).run(tmp_path / "undecided", force=True)
    for choice, expected in (("points_table", "60"), ("activity", "50")):
        settings = {**SETTINGS, "points_determinations": {"Week 1 Assignment: Report": choice}}
        Builder(model, settings, REPO_ROOT).run(tmp_path / choice, force=True)
        dropbox = (tmp_path / choice / "package" / "dropbox_d2l.xml").read_text(encoding="utf-8")
        assert f'out_of="{expected}.000000000"' in dropbox
        assert verify_build(tmp_path / choice, model) == 0


def test_bracket_tags_mark_dropdowns_and_callouts(tmp_path):
    overview = ["Page: Week 1 Overview", "Page heading: Week 1 – Sample Topic", "Overview paragraph one.",
                "[dropdown: First Idea]", "Inside the first panel.", "[/dropdown]",
                "[callout]", "Inside the callout.", "[/callout]", "Closing overview paragraph."]
    model = model_with_joins(synth.canonical(tmp_path / "brackets.docx", overview=overview))
    assert errors(model) == []
    blocks = model["weeks"][0]["overview_page"]["blocks"]
    dropdown = next(b for b in blocks if b["kind"] == "dropdown")
    assert dropdown["title"] == "First Idea" and "Inside the first panel." in "".join(r["text"] for b in dropdown["blocks"] for r in b["runs"])
    callout = next(b for b in blocks if b["kind"] == "visual")
    assert callout["cue"].lower().startswith("callout")
    assert V.parse_directive("[ /Dropdown ]") == (V.D_DROPDOWN_END, "")
    assert V.parse_directive("[timeline]") == (V.D_VISUAL, "Timeline")
    assert V.parse_directive("[not a tag]") is None


def test_a_module_row_before_week_1_holds_its_own_objects(tmp_path):
    pre_week = [
        ["Getting Started"], ["Module: Getting Started", "Page: Orientation", "Orientation text.",
                              "Discussion: Orientation Questions", "Forum: Community Forum", "Ask about getting started here."],
        ["Staff Notes"], ["Module: Staff Notes", "Visibility: Hidden from students",
                          "Link: Help Desk", "URL: https://example.org/help", "Visibility: Hidden from students"],
    ]
    source = synth.canonical(tmp_path / "modules.docx", pre_week=pre_week)
    model = model_with_joins(source)
    assert errors(model) == []
    page = next(p for p in model["pages"] if p["title"] == "Orientation")
    assert page["placement"] == ["Getting Started"] and page["placement_source"] == "module row"
    link = next(l for l in model["links"] if l["title"] == "Help Desk")
    assert link["placement"] == ["Staff Notes"]
    discussion = next(a for a in model["course_activities"] if a["title"] == "Orientation Questions")
    assert discussion["placement"] == ["Getting Started"]
    placed = {j["from"] for j in model["joins"] if j["kind"] == "object→module" and j["to"] == "Getting Started"}
    assert {"Page: Orientation", "Discussion: Orientation Questions"} <= placed


@pytest.mark.skipif(not HAS_WORKBENCH, reason="Workbench building blocks not present")
def test_internal_link_brackets_become_links_or_visible_errors(tmp_path):
    from blueprint_packager.build import Builder

    overview = ["Page: Week 1 Overview", "Page heading: Week 1 – Sample Topic",
                "Start with the [link to assignment: Week 1 Assignment: Report | report] and read [link to page: Welcome to SYN 101].",
                "A broken one: [link to discussion: Nowhere]."]
    body = [("The Report", "Heading 3"), "See the [link to page: Welcome to SYN 101] before you start."]
    model = model_with_joins(synth.canonical(tmp_path / "links.docx", overview=overview, assignment_body=body))
    assert errors(model) == []
    result = Builder(model, SETTINGS, REPO_ROOT).run(tmp_path / "out", force=True)
    page = (tmp_path / "out" / "package" / "course_pages" / "week-1-overview.html").read_text(encoding="utf-8")
    assert '<a href="welcome-to-syn-101.html">Welcome to SYN 101</a>' in page
    assert "type=dropbox&amp;rCode=" in page and ">report</a>" in page
    assert "［link to discussion: Nowhere］" in page  # left visible, in full-width brackets, so a reader sees it
    codes = {n["code"] for n in result["errors"]}
    assert {"LINK_TARGET_UNKNOWN", "LINK_TARGET_UNREACHABLE"} <= codes  # the broken link, and a page link inside an activity



# ------------------------------------------------------------ heading convention (one to one)
def _styles(docx_path):
    return [(p.text().strip(), p.style.title()) for p in read_docx(docx_path).paras() if p.heading_level]


def test_normalizer_writes_the_frame_at_heading_1_and_shifts_an_older_source(tmp_path):
    source = synth.messy(tmp_path / "messy.docx")  # an older source: frame at Heading 2, content from Heading 3
    map_path = tmp_path / "map.json"
    map_path.write_text(json.dumps(course_map_for(source)), encoding="utf-8")
    out = tmp_path / "normalized.docx"
    result = normalize(source, map_path, out)
    assert result["heading_convention"] == {"source_frame_level": 2, "frame_level": 1, "frame_shift": 1}
    styles = dict(_styles(out))
    assert styles[V.H_COURSE_INTRO] == "Heading 1" and styles["Week 1 - Sample Topic"] == "Heading 1"
    assert styles["The Report"] == "Heading 2"  # map said "heading 3" in the source's terms
    assert styles["Report Rubric"] == "Heading 2"  # rubric names sit one under the frame
    assert prove(result, out)["passed"]
    model = read_model(out)
    assert model["conventions"] == {"frame_level": 1, "content_top": 2, "heading_offset": 0}
    assert not [d for d in model["diagnostics"] if d["code"] == "FRAME_LEVEL_LEGACY"]


def test_reader_accepts_an_older_frame_with_an_offset_and_pages_still_map_one_to_one(tmp_path):
    new = read_model(synth.canonical(tmp_path / "new.docx", frame_level=1))
    old = read_model(synth.canonical(tmp_path / "old.docx", frame_level=2))
    assert new["conventions"]["heading_offset"] == 0 and old["conventions"]["heading_offset"] == 1
    assert [d["code"] for d in old["diagnostics"] if d["code"] == "FRAME_LEVEL_LEGACY"] == ["FRAME_LEVEL_LEGACY"]
    # the same structure is read from both
    strip = lambda m: {k: v for k, v in m.items() if k not in ("diagnostics", "conventions", "source", "reader")}  # noqa: E731
    new_titles = [w["overview_page"]["title"] for w in new["weeks"]]
    assert new_titles == [w["overview_page"]["title"] for w in old["weeks"]]
    assert [b["level"] for b in new["weeks"][0]["overview_page"]["blocks"] if b["kind"] == "heading"] == [2]
    assert [b["level"] for b in old["weeks"][0]["overview_page"]["blocks"] if b["kind"] == "heading"] == [3]


@pytest.mark.skipif(not HAS_WORKBENCH, reason="Workbench building blocks not present")
def test_pages_render_word_levels_one_to_one(tmp_path):
    from blueprint_packager.build import Builder

    overview = ["Page: Week 1 Overview", "Page heading: Week 1 – Sample Topic", "Overview paragraph one.",
                ("Key Ideas", "Heading 2"), "A paragraph.", ("A finer point", "Heading 3"), "Another.", ("Finer still", "Heading 4"), "Last."]
    body = [("The Report", "Heading 2"), "Write it.", ("Submission Requirements", "Heading 3"), "One document.", ("Format", "Heading 4"), "PDF."]
    pages = {}
    for label, level in (("new", 1), ("old", 2)):
        model = model_with_joins(synth.canonical(tmp_path / f"{label}.docx", overview=overview, assignment_body=body, frame_level=level))
        Builder(model, SETTINGS, REPO_ROOT).run(tmp_path / label, force=True)
        pages[label] = (tmp_path / label / "package" / "course_pages" / "week-1-overview.html").read_text(encoding="utf-8")
    for page in pages.values():
        assert page.count("<h1>") == 1
        assert "<h2>Key Ideas</h2>" in page and "<h3>A finer point</h3>" in page and "<h4>Finer still</h4>" in page
    # the older frame renders the same page: levels are read one up, so nothing changes on the page
    assert pages["new"] == pages["old"]


# ------------------------------------------------------------ editor fields, sublines, bullets, re-export forms
@pytest.mark.skipif(not HAS_WORKBENCH, reason="Workbench building blocks not present")
def test_editor_fields_carry_the_templates_modest_styles_unless_the_build_asks_for_plain(tmp_path):
    import html as html_mod
    import re

    from blueprint_packager.build import Builder

    body = [("The Report", "Heading 2"), "Write it.", ("1. First step\nWhat do you notice?", "Heading 3"), "One document."]
    model = model_with_joins(synth.canonical(tmp_path / "plain.docx", assignment_body=body))

    def description(settings, out):
        Builder(model, settings, REPO_ROOT).run(tmp_path / out, force=True)
        xml = (tmp_path / out / "package" / "dropbox_d2l.xml").read_text(encoding="utf-8")
        return html_mod.unescape(re.search(r"<instructions[^>]*>(.*?)</instructions>|<description[^>]*>(.*?)</description>", xml, re.S).group(0))

    plain = description({**SETTINGS, "fragment_styles": "plain"}, "plain")
    assert 'style="' not in plain and "<div" not in plain  # the tenant's stylesheet does the styling
    assert "<h3>1. First step<br><small>What do you notice?</small></h3>" in plain  # the author's second line is a subline
    styled = description(SETTINGS, "styled")  # the default since 2026-10-08: the template's modest set
    assert 'style="' in styled and "<text><div style=" in styled


def test_bullet_operation_makes_paragraphs_list_items_like_another(tmp_path):
    source = synth.messy(tmp_path / "messy.docx")
    cmap = course_map_for(source)
    pid = lambda text: _pid(source, text)  # noqa: E731
    from docx import Document as _D

    d = _D(str(source))  # the messy fixture has no list; add one paragraph to copy the list definition from
    synth.make_list_item(d.add_paragraph("A bullet to copy from."))
    d.save(str(source))
    cmap = course_map_for(source)
    doc = read_docx(source)
    like = next(p.id for p in doc.paras() if p.list_num_id is not None)
    cmap["notes"]["content"].append(like)  # account for the added paragraph so the fidelity proof still closes
    target = pid("Write a short synthetic report.")
    cmap["bullet"] = [{"ids": [target], "like": like, "reason": "test"}]
    map_path = tmp_path / "map.json"
    map_path.write_text(json.dumps(cmap), encoding="utf-8")
    out = tmp_path / "normalized.docx"
    result = normalize(source, map_path, out)
    assert prove(result, out)["passed"]
    written = next(p for p in read_docx(out).paras() if p.text().strip() == "Write a short synthetic report.")
    assert written.list_num_id is not None
    assert result["dispositions"][target].disposition == "kept_restyled"


def test_verifier_tolerates_re_export_forms():
    from blueprint_packager.verify import same_instant, straight_quotes

    assert same_instant("2026-11-09T04:59:00.000Z") == same_instant("2026-11-09T04:59:00") == "2026-11-09T04:59:00"
    assert straight_quotes("Week 5 The “Really?” Partnership Test") == 'Week 5 The "Really?" Partnership Test'


def test_the_default_page_template_is_the_composite():
    from blueprint_packager.html_render import Template

    default = Template.load("composite-navy")
    assert default.id == "composite-navy"
    assert default.manifest["surfaces"]["page"]["heading_map"] == {"2": "h2", "3": "h3", "4": "h4", "5": "h5", "6": "h6"}
    assert {"h2", "p", "callout", "callout_gold"} <= set(default.manifest["surfaces"]["fragment"]["styles"])  # the modest set for editor fields
    assert set(default.manifest["callout_tones"]) == {"default", "gold", "grey"}



# ------------------------------------------------------------ the blank blueprint
def test_a_blank_blueprint_reads_clean_with_a_checklist_row_in_every_week(tmp_path):
    from blueprint_packager.blank import write_blank

    path = write_blank(tmp_path / "blank.docx", course_number="NEW 101", course_title="Course Title", term="Spring A 2027",
                       weeks=3, topics=["Foundations", "Second Topic"])
    model = model_with_joins(path)
    assert errors(model) == []
    assert model["conventions"]["frame_level"] == 1
    assert [w["title"] for w in model["weeks"]] == ["Week 1 - Foundations", "Week 2 - Second Topic", "Week 3 - (topic)"]
    for week in model["weeks"]:
        assert week["checklist"]["derive"] is True  # the Items line is pre-filled in every week
        assert week["overview_page"]["page_heading"].startswith(f"Week {week['number']} – ")
        assert len(week["assignments"]) == 1 and len(week["discussions"]) == 1
    assert sorted(r["name"] for r in model["rubrics"]) == ["Standard Assignment Rubric", "Standard Discussion Rubric"]


@pytest.mark.skipif(not HAS_WORKBENCH, reason="Workbench building blocks not present")
def test_a_blank_blueprint_builds_a_package(tmp_path):
    from blueprint_packager.blank import write_blank
    from blueprint_packager.build import Builder

    model = model_with_joins(write_blank(tmp_path / "blank.docx", course_number="NEW 101", course_title="Course Title", term="Spring A 2027", weeks=2))
    result = Builder(model, SETTINGS, REPO_ROOT).run(tmp_path / "out", force=True)
    assert result["errors"] == []
    names = set(zipfile.ZipFile(result["zip"]).namelist())
    assert {"imsmanifest.xml", "grades_d2l.xml", "rubrics_d2l.xml", "dropbox_d2l.xml", "checklist_d2l.xml"} <= names


# ------------------------------------------------------------ callout tones, wrap, bullet guard
@pytest.mark.skipif(not HAS_WORKBENCH, reason="Workbench building blocks not present")
def test_callout_tones_and_colours_from_bracket_attributes(tmp_path):
    from blueprint_packager.build import Builder

    overview = ["Page: Week 1 Overview", "Page heading: Week 1 – Sample Topic", "Lead paragraph.",
                "[callout tone=gold]", "Keep this distinction in mind:", "[/callout]",
                "[callout color=#336699]", "A custom colour.", "[/callout]",
                "[callout tone=plaid]", "Unknown tone.", "[/callout]"]
    model = model_with_joins(synth.canonical(tmp_path / "tones.docx", overview=overview, frame_level=1))
    assert errors(model) == []
    settings = {**SETTINGS, "template": "composite-navy"}
    result = Builder(model, settings, REPO_ROOT).run(tmp_path / "out", force=True)
    page = (tmp_path / "out" / "package" / "course_pages" / "week-1-overview.html").read_text(encoding="utf-8")
    assert '<div class="callout callout--gold">' in page
    assert '<div class="callout" style="border-left-color: #336699;">' in page
    assert page.count('<div class="callout">') == 1  # the unknown tone fell back to the default
    notes = json.loads((tmp_path / "out" / "object_map.json").read_text(encoding="utf-8"))["notes"]
    assert any(n["code"] == "CALLOUT_TONE_UNKNOWN" for n in notes)
    assert '<a class="skip" href="#main">' in page and 'loading="lazy"' not in page  # no video on this page


def test_wrap_writes_bracket_lines_around_kept_paragraphs_and_bullet_never_doubles(tmp_path):
    source = synth.messy(tmp_path / "messy.docx")
    cmap = course_map_for(source)
    first = _pid(source, "Write a short synthetic report.")
    cmap["wrap"] = [{"first": first, "last": first, "block": "callout", "attrs": {"tone": "gold"}, "reason": "test"}]
    from docx import Document as _D

    d = _D(str(source))
    synth.make_list_item(d.add_paragraph("Already a bullet."))
    d.save(str(source))
    cmap = course_map_for(source)
    cmap["wrap"] = [{"first": first, "last": first, "block": "callout", "attrs": {"tone": "gold"}, "reason": "test"}]
    doc = read_docx(source)
    already = next(p.id for p in doc.paras() if p.list_num_id is not None)
    cmap["notes"]["content"].append(already)
    cmap["bullet"] = [{"ids": [already], "like": already, "reason": "would double-bullet"}]
    map_path = tmp_path / "map.json"
    map_path.write_text(json.dumps(cmap), encoding="utf-8")
    out = tmp_path / "normalized.docx"
    result = normalize(source, map_path, out)
    assert prove(result, out)["passed"]
    texts = [p.text().strip() for p in read_docx(out).paras()]
    i = texts.index("Write a short synthetic report.")
    assert texts[i - 1] == "[callout tone=gold]" and texts[i + 1] == "[/callout]"
    assert "already a list item; bullet operation skipped" in (result["records"][[r.origin for r in result["records"]].index(already)].note)
    model = read_model(out)
    callouts = [b for b in model["weeks"][0]["assignments"][0]["blocks"] if b["kind"] == "visual"]
    assert callouts and callouts[0]["attrs"] == {"tone": "gold"}

