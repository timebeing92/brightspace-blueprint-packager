"""Synthetic closed-cycle and failure cases for the explicit Wizard adapter."""
import copy
import hashlib
import json
import os
import subprocess
import sys
from pathlib import Path

import pytest
import synth

from blueprint_packager.build import Builder
from blueprint_packager.joins import resolve
from blueprint_packager.reader import read_model
from blueprint_packager import workbench as WB
from blueprint_packager.wizard_authoring import convert_bundle

SETTINGS = {
    "schema": "blueprint_packager.build_settings/0.1",
    "package": {"title": "SYN 101 conversion test", "manifest_identifier": "D2L_SYN101",
                "base_resource_code": "00000000-0000-4000-8000-000000000101"},
    "template": "bumk515-classic",
    "forum_order": ["Community Forum", "Discussions"],
}


def joined(path):
    model = read_model(path)
    result = resolve(model)
    model["joins"] = result["joins"]
    model["diagnostics"].extend(result["diagnostics"])
    return model


def extract(package, directory):
    command = [sys.executable, "-B", str(WB.scripts_dir() / "build_blueprint_bundle.py"), str(package),
               "--label", "synthetic", "--output-dir", str(directory), "--no-docx", "--no-syllabus-fetch", "--quiet"]
    process = subprocess.run(command, capture_output=True, text=True,
                             env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1"})
    assert process.returncode == 0, process.stdout + process.stderr
    return directory / "synthetic__blueprint_bundle"


@pytest.fixture(scope="module")
def cycle(tmp_path_factory):
    root = tmp_path_factory.mktemp("wizard-cycle")
    model = joined(synth.canonical(root / "source.docx"))
    week = model["weeks"][0]
    for key, kind in (("assignments", "Assignment"), ("discussions", "Discussion")):
        activity = copy.deepcopy(week[key][0])
        activity["title"] = activity["grade_item"] = f"Week 1 {kind}: Second task"
        activity["points"] = 20.
        activity["blocks"] = [{"kind": "p", "runs": [{"text": f"Complete the second synthetic {kind.lower()}."}]}]
        week[key].append(activity)
        model["grade_plan"]["items"].append({"name": activity["title"], "points": 20.})
        rubric = next(r for r in model["rubrics"] if r["name"] == activity["rubric"])
        rubric["used_by"].append(f"{kind}: {activity['title']}")
    model["grade_plan"].update(total_stated=140., total_computed=140.)
    week["overview_page"]["blocks"].extend([
        {"kind": "video", "title": "Synthetic overview video", "href": "https://www.youtube.com/embed/synthetic01"},
        {"kind": "p", "runs": [{"text": "A bold phrase", "b": True}, {"text": " and "},
                                 {"text": "an italic phrase", "i": True},
                                 {"text": " with a reference.", "href": "https://example.org/reference"}]},
        {"kind": "p", "runs": [{"text": "Read [link to assignment: Week 1 Assignment: Report | the report task]."}]},
        {"kind": "p", "runs": [{"text": "Return to [link to page: Welcome to SYN 101 | the welcome page]."}]},
    ])
    joins = resolve(model)
    model["joins"] = joins["joins"]
    model["diagnostics"] = joins["diagnostics"]
    assert not [d for d in model["diagnostics"] if d["severity"] == "error"]
    built = Builder(model, SETTINGS, WB.REPO_DIR).run(root / "first-build", force=False)
    assert not built["errors"]
    bundle = extract(built["zip"], root / "first-extraction")
    return root, model, bundle


def semantic(blocks):
    """Compare authored text/structure, excluding extraction-only style evidence."""
    result = []
    for block in blocks:
        kind = block["kind"]
        entry = {"kind": kind}
        for key in ("title", "cue", "href", "level", "ordered"):
            if key in block:
                entry[key] = block[key]
        if "runs" in block:
            entry["runs"] = [{k: v for k, v in r.items() if k != "b" or kind != "heading"} for r in block["runs"]]
        if "blocks" in block:
            entry["blocks"] = semantic(block["blocks"])
        result.append(entry)
    return result


def test_closed_cycle_preserves_multiple_activities_media_and_joins(cycle):
    root, source, bundle = cycle
    before_hashes = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in bundle.glob("*.json")}
    output = root / "returned.docx"
    report = convert_bundle(bundle, output)
    assert report["status"] == "converted", report["diagnostics"]
    returned = joined(output)
    assert not [d for d in returned["diagnostics"] if d["severity"] == "error"]
    assert {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in bundle.glob("*.json")} == before_hashes
    assert len(report["sources"]) == 4
    assert all(len(s["sha256"]) == 64 for s in report["sources"])
    assert returned["grade_plan"]["items"] == source["grade_plan"]["items"]
    assert returned["grade_plan"]["total_computed"] == 140
    for key in ("pages", "modules", "links", "course_activities"):
        assert [(o["type"], o["title"], o.get("placement"), o.get("visibility")) for o in returned[key]] == [
            (o["type"], o["title"], o.get("placement"), o.get("visibility")) for o in source[key]]
    a, b = source["weeks"][0], returned["weeks"][0]
    assert (a["title"], a["overview_page"]["title"], a["overview_page"]["page_heading"]) == (
        b["title"], b["overview_page"]["title"], b["overview_page"]["page_heading"])
    assert semantic(a["overview_page"]["blocks"]) == semantic(b["overview_page"]["blocks"])
    assert semantic(a["objectives"]) == semantic(b["objectives"])
    assert semantic(a["resources"]) == semantic(b["resources"])
    for key in ("assignments", "discussions"):
        assert len(b[key]) == 2
        for left, right in zip(a[key], b[key]):
            for attr in ("title", "points", "grade_item", "rubric", "forum"):
                assert left.get(attr) == right.get(attr)
            assert semantic(left["blocks"]) == semantic(right["blocks"])
    assert [[r["text"] for r in item["runs"]] for item in a["checklist"]["items"]] == [
        [r["text"] for r in item["runs"]] for item in b["checklist"]["items"]]
    for left, right in zip(source["rubrics"], returned["rubrics"]):
        assert left["name"] == right["name"]
        assert left["used_by"] == right["used_by"]
        assert [(c["name"], c["weight"], semantic(c["cells"][0])) for c in left["grid"]["criteria"]] == [
            (c["name"], c["weight"], semantic(c["cells"][0])) for c in right["grid"]["criteria"]]
    rebuilt = Builder(returned, SETTINGS, WB.REPO_DIR).run(root / "second-build", force=False)
    assert not rebuilt["errors"]
    second_bundle = extract(rebuilt["zip"], root / "second-extraction")
    # Extracted HTML content and object payloads should be stable through the
    # second generation even though source-file provenance comments differ.
    for suffix in ("__blueprint.json", "__rubrics.json"):
        first = json.loads(next(bundle.glob("*" + suffix)).read_text(encoding="utf-8"))
        second = json.loads(next(second_bundle.glob("*" + suffix)).read_text(encoding="utf-8"))
        if suffix == "__blueprint.json":
            for key in ("front_matter", "before_week_1", "weeks", "unplaced_activities"):
                assert first[key] == second[key]
        else:
            assert first["rubrics"] == second["rubrics"]


def mutate_bundle(cycle, tmp_path, role, change):
    import shutil
    source = cycle[2]
    copied = tmp_path / "bundle"
    shutil.copytree(source, copied)
    suffix = {"blueprint": "__blueprint.json", "activities": "__course_activities.json",
              "structure": "__course_structure.json", "rubrics": "__rubrics.json"}[role]
    path = next(copied.glob("*" + suffix))
    payload = json.loads(path.read_text(encoding="utf-8"))  # the bundle is UTF-8; Windows would otherwise read it as cp1252
    change(payload)
    path.write_text(json.dumps(payload), encoding="utf-8")
    return copied


@pytest.mark.parametrize(("role", "change", "code"), [
    ("activities", lambda x: x["dropbox_folders"][0].update(date_due="2026-10-09T23:59:00Z"), "ACTIVITY_SETTINGS_UNSUPPORTED"),
    ("activities", lambda x: x["dropbox_folders"][0].update(grade_item_code="missing"), "GRADE_JOIN_UNRESOLVED"),
    ("activities", lambda x: x["dropbox_folders"][0]["instructions_blocks"].append({"kind": "practice", "runs": [{"text": "Practice"}]}), "BLOCK_UNSUPPORTED"),
    ("structure", lambda x: x.update(run_id="different-extraction"), "BUNDLE_RUN_MISMATCH"),
    ("structure", lambda x: x["tree"][-1].update(is_hidden=True), "WEEK_VISIBILITY_UNSUPPORTED"),
    ("structure", lambda x: x["tree"][-1]["children"][-1].update(title="Different manifest title"), "MANIFEST_TITLE_DIFFERS"),
    ("rubrics", lambda x: x["rubrics"][0]["criteria"][0]["cells"][0].update(points=99), "RUBRIC_SCORE_UNREPRESENTABLE"),
    ("blueprint", lambda x: x["front_matter"]["course_introduction"].append({"kind": "p", "runs": [{"text": "Unaccounted authored passage."}]}), "BLUEPRINT_CONTENT_UNACCOUNTED"),
])
def test_unsupported_or_inconsistent_evidence_blocks_conversion(cycle, tmp_path, role, change, code):
    bundle = mutate_bundle(cycle, tmp_path, role, change)
    output = tmp_path / "blocked.docx"
    result = convert_bundle(bundle, output)
    assert result["status"] == "blocked"
    assert code in {d["code"] for d in result["diagnostics"]}
    assert not output.exists()
    assert output.with_suffix(".conversion.json").is_file()


def test_refuses_overwrite_and_writes_outside_evidence(cycle, tmp_path):
    output = tmp_path / "authored.docx"
    output.write_bytes(b"Existing authored document")
    with pytest.raises(FileExistsError):
        convert_bundle(cycle[2], output)
    assert output.read_bytes() == b"Existing authored document"
    with pytest.raises(ValueError, match="immutable"):
        convert_bundle(cycle[2], cycle[2] / "output.docx")


def test_missing_or_unknown_contracts_produce_reviewable_diagnostics(cycle, tmp_path):
    bundle = mutate_bundle(cycle, tmp_path, "blueprint", lambda x: x.update(schema="coursecraft.blueprint/999"))
    result = convert_bundle(bundle, tmp_path / "unknown.docx")
    assert result["status"] == "blocked"
    assert result["diagnostics"][0]["code"] == "BUNDLE_SCHEMA"
    missing = tmp_path / "empty"; missing.mkdir()
    result = convert_bundle(missing, tmp_path / "missing.docx")
    assert len(result["diagnostics"]) == 4


def test_extracted_explicit_callout_tone_is_carried_into_authoring(cycle, tmp_path):
    def add_tone(payload):
        block = next(b for b in payload["weeks"][0]["overview"] if b["kind"] == "visual")
        block["meta"]["tone"] = "gold"
    bundle = mutate_bundle(cycle, tmp_path, "blueprint", add_tone)
    output = tmp_path / "tone.docx"
    result = convert_bundle(bundle, output)
    assert result["status"] == "converted", result["diagnostics"]
    blocks = read_model(output)["weeks"][0]["overview_page"]["blocks"]
    assert next(b for b in blocks if b["kind"] == "visual")["attrs"] == {"tone": "gold"}


def test_no_rubric_package_converts_using_explicit_inventory_evidence(cycle, tmp_path):
    model = copy.deepcopy(cycle[1])
    model["rubrics"] = []
    for week in model["weeks"]:
        for key in ("assignments", "discussions"):
            for activity in week[key]:
                activity.pop("rubric", None)
    resolved = resolve(model)
    model["joins"], model["diagnostics"] = resolved["joins"], resolved["diagnostics"]
    first = Builder(model, SETTINGS, WB.REPO_DIR).run(tmp_path / "without-rubric", force=False)
    bundle = extract(first["zip"], tmp_path / "extraction")
    assert not list(bundle.glob("*__rubrics.json"))
    output = tmp_path / "no-rubric.docx"
    result = convert_bundle(bundle, output)
    assert result["status"] == "converted", result["diagnostics"]
    assert any(s["role"] == "inventory" and s["sha256"] for s in result["sources"])
    assert any(e["source"] == "rubrics" and e["disposition"] == "not_applicable" for e in result["ledger"])
    returned = joined(output)
    assert returned["rubrics"] == []
    rebuilt = Builder(returned, SETTINGS, WB.REPO_DIR).run(tmp_path / "rebuilt", force=False)
    assert not rebuilt["errors"]
    second = extract(rebuilt["zip"], tmp_path / "second-extraction")
    assert not list(second.glob("*__rubrics.json"))
    assert json.loads(next(bundle.glob("*__blueprint.json")).read_text(encoding="utf-8"))["weeks"] == json.loads(
        next(second.glob("*__blueprint.json")).read_text(encoding="utf-8"))["weeks"]


def test_absent_rubric_contract_with_existing_payload_is_an_error(cycle, tmp_path):
    import shutil
    bundle = tmp_path / "bundle"
    shutil.copytree(cycle[2], bundle)
    next(bundle.glob("*__rubrics.json")).unlink()
    result = convert_bundle(bundle, tmp_path / "missing-rubric.docx")
    assert result["status"] == "blocked"
    assert "RUBRIC_CONTRACT_MISSING" in {d["code"] for d in result["diagnostics"]}


@pytest.mark.parametrize("conflict", ["document", "ledger"])
def test_exclusive_publication_preserves_a_concurrent_file(cycle, tmp_path, monkeypatch, conflict):
    from blueprint_packager import wizard_authoring as adapter
    output = tmp_path / "concurrent.docx"
    ledger = output.with_suffix(".conversion.json")
    target = output if conflict == "document" else ledger
    real_link = adapter.os.link
    def concurrent_link(source, destination):
        if Path(destination) == target:
            target.write_bytes(b"Another writer's completed artifact")
        return real_link(source, destination)
    monkeypatch.setattr(adapter.os, "link", concurrent_link)
    with pytest.raises(FileExistsError):
        convert_bundle(cycle[2], output)
    assert target.read_bytes() == b"Another writer's completed artifact"
    assert not (ledger if conflict == "document" else output).exists()
    assert not list(tmp_path.glob(".wizard-*"))


# --------------------------------------------------------------------------- #
# Blueprint Wizard compatibility (2026-10-13): tables, panel labels, heading ownership
# --------------------------------------------------------------------------- #
from docx import Document as _Document  # noqa: E402
from docx.oxml.ns import qn as _qn  # noqa: E402

from blueprint_packager import vocabulary as V  # noqa: E402
from blueprint_packager import wizard_authoring as adapter  # noqa: E402


def _table_block(header_rows):
    cell = lambda text: [{"text": text, "href": ""}]  # noqa: E731
    return {"kind": "table", "level": 0, "runs": [], "header_rows": header_rows,
            "rows": [[cell("Column A"), cell("Column B")], [cell("a1"), cell("b1")], [cell("a2"), cell("b2")]]}


def _panel(label, children, meta=None):
    block = {"kind": "visual", "level": 0, "runs": [{"text": label, "href": ""}], "blocks": children}
    if meta:
        block["meta"] = meta
    return block


def _heading(text, level=2):
    return {"kind": "label", "level": level, "runs": [{"text": text, "href": ""}]}


def _para(text):
    return {"kind": "p", "level": 0, "runs": [{"text": text, "href": ""}]}


def _first_table_in_week_cell(path):
    """The first table nested inside a cell of a week table (the grade plan sits in the body)."""
    for table in _Document(str(path)).tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.tables:
                    return cell.tables[0]
    raise AssertionError("no table inside a week cell")


@pytest.mark.parametrize("header_rows", [0, 1])
def test_extracted_tables_convert_and_keep_only_declared_header_rows(cycle, tmp_path, header_rows):
    """The candidate crashed on `_Cell.add_table(width=)` and shaded a first row the
    sandbox tables (header_rows: 0) never declared."""
    bundle = mutate_bundle(cycle, tmp_path, "blueprint",
                           lambda x: x["weeks"][0]["overview"].append(_table_block(header_rows)))
    output = tmp_path / f"table{header_rows}.docx"
    result = convert_bundle(bundle, output)
    assert result["status"] == "converted", result["diagnostics"]
    table = _first_table_in_week_cell(output)
    first_row = table.rows[0]
    shaded = [c._tc.tcPr is not None and c._tc.tcPr.find(_qn("w:shd")) is not None for c in first_row.cells]
    repeated = first_row._tr.trPr is not None and first_row._tr.trPr.find(_qn("w:tblHeader")) is not None
    bold = [all(r.bold for r in c.paragraphs[0].runs) for c in first_row.cells]
    assert (all(shaded), repeated, all(bold)) == ((header_rows > 0),) * 3
    read_back = read_model(output)["weeks"][0]["overview_page"]["blocks"]
    tables = [b for b in read_back if b["kind"] == "table"]
    assert tables and [[r["text"] for b in cell for r in b["runs"]] for cell in tables[-1]["rows"][1]] == [["a1"], ["b1"]]


@pytest.mark.parametrize(("label", "meta", "opener", "cue", "attrs"), [
    ("Project callout", {"tone": "gold"}, "[callout tone=gold]", "Callout", {"tone": "gold"}),
    ("Panel", {"color": "#336699"}, "[box color=#336699: Panel]", "Styled section", {"color": "#336699"}),
    ("Accordion", {}, "[box: Accordion]", "Styled section", {}),
    ("Discussion note", {}, "[callout]", "Callout", {}),
    ("Timeline", {}, "[timeline]", "Timeline", {}),
    ("Framed section", {}, "[box: Framed section]", "Styled section", {}),
])
def test_extracted_panel_labels_normalise_to_the_bracket_vocabulary(cycle, tmp_path, label, meta, opener, cue, attrs):
    """Attributed labels other than the three exact strings used to be refused."""
    bundle = mutate_bundle(cycle, tmp_path, "blueprint",
                           lambda x: x["weeks"][0]["overview"].append(_panel(label, [_para("Inside the panel.")], meta)))
    output = tmp_path / "labels.docx"
    result = convert_bundle(bundle, output)
    assert result["status"] == "converted", result["diagnostics"]
    assert not [d for d in result["diagnostics"] if d["code"] == "VISUAL_ATTRIBUTES_UNSUPPORTED"]
    texts = [p.text for p in _Document(str(output)).paragraphs] + [
        p.text for t in _Document(str(output)).tables for r in t.rows for c in r.cells for p in c.paragraphs]
    assert opener in texts
    assert V.parse_bracket_block(opener) == (V.D_VISUAL, cue) and V.parse_bracket_attrs(opener) == attrs
    blocks = read_model(output)["weeks"][0]["overview_page"]["blocks"]
    panel = [b for b in blocks if b["kind"] == "visual"][-1]
    assert (panel["cue"], panel.get("attrs") or {}) == (cue, attrs)


def test_bracket_keyword_mirrors_the_wizard():
    assert [V.bracket_keyword(c) for c in ("Callout", "Project callout", "Note", "Discussion note", "Timeline", "Accordion", "Styled section", "")] == [
        "callout", "callout", "callout", "callout", "timeline", "box", "box", "box"]
    assert V.bracket_opener("Accordion", {"tone": "gold"}) == "[box tone=gold: Accordion]"
    assert V.bracket_opener("Callout", {"tone": "", "color": "#abcdef"}) == "[callout color=#abcdef]"


def test_a_heading_a_panel_owns_is_written_once(cycle, tmp_path):
    """A structure segment marked heading_owner: "panel" and a resources section
    whose first block is a panel carrying its label each print the heading once."""
    def own_in_structure(payload):
        page = next(p for p in payload["html_topics"] if p.get("body_segments"))
        page["body_segments"].append({
            "heading": "Owned Heading", "level": 2, "heading_owner": "panel",
            "blocks": [_panel("Callout", [_heading("Owned Heading"), _para("Owned body.")])],
            "text": "Owned Heading Owned body.",
        })
        page["body_segments"].append({
            "heading": "Plain Heading", "level": 2,
            "blocks": [_para("Plain body.")], "text": "Plain Heading Plain body.",
        })
        page["body_text"] = page.get("body_text", "") + " Owned Heading Owned body. Plain Heading Plain body."
    bundle = mutate_bundle(cycle, tmp_path, "structure", own_in_structure)
    def own_in_resources(payload):
        payload["weeks"][0]["resources"].append({
            "label": "Reading Panel", "level": 2, "source_page": "Week 1 Overview",
            "blocks": [_panel("Resources section", [_heading("Reading Panel"), _para("Read this.")])],
        })
    import json as _json
    blueprint_path = next(bundle.glob("*__blueprint.json"))
    payload = _json.loads(blueprint_path.read_text(encoding="utf-8"))
    own_in_resources(payload)
    blueprint_path.write_text(_json.dumps(payload), encoding="utf-8")
    output = tmp_path / "owned.docx"
    result = convert_bundle(bundle, output)
    assert result["status"] in {"converted", "blocked"}, result["diagnostics"]
    doc = _Document(str(output)) if output.exists() else None
    if doc is None:  # the synthetic page lives outside the overview join; the DOCX still carries it when converted
        pytest.fail(str(result["diagnostics"]))
    texts = [p.text for p in doc.paragraphs] + [p.text for t in doc.tables for r in t.rows for c in r.cells for p in c.paragraphs]
    assert texts.count("Owned Heading") == 1 and texts.count("Plain Heading") == 1
    assert texts.count("Reading Panel") == 1
    assert texts[texts.index("Reading Panel") - 1] == "[box: Resources section]"
    ledger = [(e["disposition"], e["detail"]) for e in result["ledger"] if e["disposition"] == "heading_owned_by_panel"]
    assert ("heading_owned_by_panel", "Owned Heading") in ledger and ("heading_owned_by_panel", "Reading Panel") in ledger


OWNERSHIP_FIXTURES = [
    ([_panel("Callout", [_heading("Assigned Reading"), _para("x")]), _para("y")], "Assigned Reading", True),
    ([_panel("Callout", [_heading("Assigned Reading"), _para("x")])], "assigned reading:", True),
    ([_panel("Callout", [_heading("Assigned Reading"), _para("x")])], "Other", False),
    ([_para("y"), _panel("Callout", [_heading("Assigned Reading")])], "Assigned Reading", False),
    ([_panel("Callout", [_para("lead"), _heading("Assigned Reading")])], "Assigned Reading", False),
    ([{"kind": "dropdown", "level": 0, "runs": [{"text": "More", "href": ""}], "blocks": [_heading("Term", 3)]}], "Term", True),
    ([_panel("Callout", [{"kind": "label", "level": 0, "runs": [{"text": "Points:", "href": ""}]}])], "Points", False),
    ([], "Assigned Reading", False),
]


@pytest.mark.parametrize(("blocks", "label", "expected"), OWNERSHIP_FIXTURES)
def test_panel_ownership_helper_matches_the_workbench(blocks, label, expected):
    """Parity with reconstruct_course_structure.panel_owns_label on the same fixtures.

    Runs against whichever Workbench the Packager resolves (vendored pin, or
    BPK_WORKBENCH_SCRIPTS for the candidate producer); skips when that copy
    predates the helper."""
    assert adapter.panel_owns_label(blocks, label) is expected
    import importlib.util
    scripts = WB.scripts_dir()
    spec = importlib.util.spec_from_file_location("rcs_parity", scripts / "reconstruct_course_structure.py")
    module = importlib.util.module_from_spec(spec)
    sys.dont_write_bytecode = True
    spec.loader.exec_module(module)
    if not hasattr(module, "panel_owns_label"):
        pytest.skip(f"Workbench at {scripts} predates panel_owns_label (2026-10-13)")
    assert module.panel_owns_label(blocks, label) is expected
